from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, ttk

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from integrations.dea.dea_engine import solve_dataset
from integrations.dea.excel_io import export_results, load_dea_workbook
from integrations.bootstrap.bootstrap_dematel import (
    build_defuzzify_func, load_fuzzy_scale, parse_fuzzy_label,
    read_expert_sheets, run_pipeline,
)


def _frame(parent, title, row, column=0, colspan=1):
    box = ttk.LabelFrame(parent, text=title, padding=10)
    box.grid(row=row, column=column, columnspan=colspan, sticky="nsew", padx=10, pady=8)
    return box


def _preview_table(parent, height=10, pack=True):
    frame = ttk.Frame(parent)
    if pack: frame.pack(fill="both", expand=True, padx=6, pady=5)
    tree = ttk.Treeview(frame, show="headings", height=height)
    ybar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
    xbar = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
    tree.grid(row=0, column=0, sticky="nsew"); ybar.grid(row=0, column=1, sticky="ns"); xbar.grid(row=1, column=0, sticky="ew")
    frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)
    return tree


def _show_df(tree, frame, max_rows=200, max_cols=40):
    tree.delete(*tree.get_children())
    if frame is None or frame.empty:
        tree["columns"] = (); return
    frame = frame.iloc[:max_rows, :max_cols]
    cols = [str(c) for c in frame.columns]
    tree["columns"] = cols
    for col in cols:
        tree.heading(col, text=col); tree.column(col, width=145, minwidth=85, stretch=True, anchor="center")
    for row in frame.itertuples(index=False, name=None):
        values = ["" if pd.isna(x) else f"{x:.5g}" if isinstance(x, (float, np.floating)) else str(x) for x in row]
        tree.insert("", "end", values=values)


def _template_path(relative_path):
    """Resolve shared templates in source and PyInstaller builds."""
    bundled = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[3]))
    return bundled / "Templates" / "Template" / relative_path


def generate_dea_demo(n_dmus=30, n_inputs=2, n_outputs=2, seed=42):
    """Create reproducible positive input/output values for a DEA practice file."""
    n_dmus, n_inputs, n_outputs, seed = int(n_dmus), int(n_inputs), int(n_outputs), int(seed)
    if not 2 <= n_dmus <= 2000: raise ValueError("DMUs must be between 2 and 2,000.")
    if not 1 <= n_inputs <= 10 or not 1 <= n_outputs <= 10: raise ValueError("Choose 1–10 input and output variables.")
    rng = np.random.default_rng(seed)
    X = rng.uniform(5.0, 100.0, size=(n_dmus, n_inputs))
    efficiency = rng.uniform(0.55, 1.0, size=n_dmus)
    weights = rng.uniform(0.25, 1.0, size=(n_inputs, n_outputs))
    production = np.maximum(X @ weights, 1e-6)
    Y = production * efficiency[:, None] * rng.uniform(0.96, 1.04, size=(n_dmus, n_outputs))
    data = pd.DataFrame({"DMU_ID": [f"DMU_{i:03d}" for i in range(1, n_dmus + 1)], "Period": [""] * n_dmus})
    inputs = [f"Input_{i}" for i in range(1, n_inputs + 1)]
    outputs = [f"Output_{i}" for i in range(1, n_outputs + 1)]
    for j, name in enumerate(inputs): data[name] = X[:, j]
    for j, name in enumerate(outputs): data[name] = Y[:, j]
    return data, inputs, outputs


def _dea_formula(model, orientation, rts):
    if model in ("CCR", "BCC", "SuperEfficiency"):
        if orientation == "Output":
            formula = ("max φ;  Σⱼ λⱼ xᵢⱼ ≤ xᵢ₀;  Σⱼ λⱼ yᵣⱼ ≥ φ yᵣ₀;  λⱼ ≥ 0;  efficiency = 1/φ*")
        else:
            formula = ("min θ;  Σⱼ λⱼ xᵢⱼ ≤ θ xᵢ₀;  Σⱼ λⱼ yᵣⱼ ≥ yᵣ₀;  λⱼ ≥ 0;  efficiency = θ*")
        if model == "SuperEfficiency": formula += "\nExclude the evaluated DMU from the reference set."
    elif model == "SBM":
        formula = "min ρ = [1 − (1/m)Σᵢ(sᵢ⁻/xᵢ₀)] / [1 + (1/s)Σᵣ(sᵣ⁺/yᵣ₀)];  Xλ+s⁻=x₀;  Yλ−s⁺=y₀"
    elif model == "Additive":
        formula = "max normalized sum of input and output slacks;  Xλ+s⁻=x₀;  Yλ−s⁺=y₀;  λ,s⁻,s⁺ ≥ 0"
    elif model == "CrossEfficiency":
        formula = "max Σᵣuᵣyᵣₖ;  Σᵢvᵢxᵢₖ=1;  Σᵣuᵣyᵣⱼ−Σᵢvᵢxᵢⱼ≤0;  cross-score=(u·yⱼ)/(v·xⱼ)"
    else:
        formula = "MPI = √[(Dₜ(xₜ₊₁,yₜ₊₁)/Dₜ(xₜ,yₜ)) × (Dₜ₊₁(xₜ₊₁,yₜ₊₁)/Dₜ₊₁(xₜ,yₜ))];  EC and TC decompose productivity change."
    scale = {"CRS": "No convexity restriction (constant returns to scale).", "VRS": "Σⱼλⱼ=1 (variable returns to scale).", "NIRS": "Σⱼλⱼ≤1 (non-increasing returns).", "NDRS": "Σⱼλⱼ≥1 (non-decreasing returns)."}.get(rts, "")
    note = f"{scale}\nA score of 1 is on the frontier for radial models. Model scores can have different scales."
    if model == "SBM": note += " SBM currently uses input orientation and requires positive inputs and outputs."
    if model == "SuperEfficiency": note += " Super-efficiency currently uses input orientation."
    if model == "CrossEfficiency": note += " Cross-efficiency currently supports CRS."
    if model == "Malmquist_MPI": note += " Requires multiple periods and common DMUs across adjacent periods."
    return formula, note


def build_dea_panels(app, pages):
    """Rebuild the full DEA import → model → result workflow as notebook tabs."""
    state = {"path": None, "bundle": None, "result": None, "generated": None}
    model = tk.StringVar(value="CCR"); orientation = tk.StringVar(value="Input"); rts = tk.StringVar(value="CRS")
    validation = tk.StringVar(value="Import an Excel workbook. Input data preview and validation appear here before analysis.")

    # Input and import preview
    page = pages["input"]; page.columnconfigure(0, weight=1)
    ttk.Label(page, text="Input Data", style="Title.TLabel").grid(row=0, column=0, sticky="w", padx=18, pady=(14, 2))
    ttk.Label(page, text="Import, review and validate the Data sheet before running a DEA model.", style="Muted.TLabel").grid(row=1, column=0, sticky="w", padx=18, pady=(0, 7))
    actions = ttk.Frame(page); actions.grid(row=2, column=0, sticky="new", padx=12, pady=4)
    path_label = ttk.Label(actions, text="No workbook selected", wraplength=680); path_label.pack(side="left", padx=8)
    preview = _preview_table(page, height=14, pack=False); preview.master.grid(row=3, column=0, sticky="nsew", padx=12, pady=4); page.rowconfigure(3, weight=1)
    def set_validation(message, kind="info"):
        validation.set(message)
        style = {"error": "Validation.Error.TLabel", "warning": "Validation.Warning.TLabel"}.get(kind, "Muted.TLabel")
        banner.configure(style=style)
    banner = ttk.Label(page, textvariable=validation, style="Muted.TLabel", wraplength=1080)
    banner.grid(row=4, column=0, sticky="ew", padx=18, pady=7)
    def import_workbook(path=None):
        if not path: path = filedialog.askopenfilename(filetypes=[("Excel workbooks", "*.xlsx *.xlsm")])
        if not path: return
        try:
            bundle = load_dea_workbook(path)
            data = pd.read_excel(path, sheet_name="Data")
            state.update(path=path, bundle=bundle, result=None)
            model.set(bundle.get("model", "CCR"));orientation.set(bundle.get("orientation", "Input"));rts.set(bundle.get("rts", "CRS"))
            path_label.configure(text=f"{Path(path).name}  •  {len(bundle['data'])} DMUs  •  {len(bundle['input_cols'])} inputs  •  {len(bundle['output_cols'])} outputs")
            _show_df(preview, data)
            warning_text = " ".join(bundle.get("warnings", []))
            if warning_text: set_validation("Workbook passed validation with a sample-size warning: " + warning_text, "warning")
            else: set_validation("Workbook passed validation. The preview shows the imported Data sheet; configured inputs and outputs are identified.")
            app.tabs.select(app.dea_pages["model"])
        except Exception as exc:
            state.update(path=path, bundle=None, result=None)
            set_validation(f"This workbook needs correction before analysis: {exc}", "error")
            app.error(f"DEA input validation failed:\n\n{exc}")
    app.action(actions, "Import Excel…", import_workbook, True).pack(side="left", padx=4)
    template = _template_path(Path("DataStudio") / "DEA" / "DEA_Input_Template.xlsx")
    def save_template():
        target = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="DEA_Input_Template.xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if target:
            try: shutil.copy2(template, target); set_validation(f"Blank Guide + Data template saved to {target}.")
            except Exception as exc: app.error(exc)
    app.action(actions, "Save blank template…", save_template).pack(side="left", padx=4)
    ttk.Label(page, text="Only the yellow input cells in the blank template need data. Keep DMU_ID unique; use Input_* and Output_* column names.", style="Muted.TLabel", wraplength=1080).grid(row=5, column=0, sticky="ew", padx=18, pady=4)

    # Model and formula preview
    page = pages["model"]; page.columnconfigure(0, weight=1); page.rowconfigure(3, weight=1)
    ttk.Label(page, text="Model & Formula", style="Title.TLabel").grid(row=0, column=0, sticky="w", padx=18, pady=(14, 2))
    ttk.Label(page, text="The selected formulation and restrictions update with your choices.", style="Muted.TLabel").grid(row=1, column=0, sticky="w", padx=18, pady=(0, 7))
    config = _frame(page, "Model configuration", 2)
    for c in range(6): config.columnconfigure(c, weight=1 if c % 2 else 0)
    model_box = orientation_box = rts_box = None
    for col, label, var, values in [(0,"Model family",model,["CCR","BCC","SBM","Additive","CrossEfficiency","SuperEfficiency","Malmquist_MPI"]),(2,"Orientation",orientation,["Input","Output"]),(4,"Returns to scale",rts,["CRS","VRS","NIRS","NDRS"])]:
        ttk.Label(config,text=label).grid(row=0,column=col,sticky="w",padx=6,pady=5)
        box=ttk.Combobox(config,textvariable=var,values=values,state="readonly",width=22);box.grid(row=0,column=col+1,sticky="ew",padx=6,pady=5)
        if col==0:model_box=box
        elif col==2:orientation_box=box
        else:rts_box=box
    formula_box = _frame(page, "Formulation preview", 3); formula_box.pack_propagate(False); formula_box.configure(height=300)
    formula_text = tk.Text(formula_box, height=8, wrap="word", font=(app.ui_font, 11), bg="#FFFFFF", relief="flat", padx=10, pady=8)
    formula_text.pack(fill="both", expand=True)
    def update_formula(*_):
        selected=model.get()
        fixed_rts={"CCR":"CRS","BCC":"VRS","CrossEfficiency":"CRS","Malmquist_MPI":"CRS"}
        if selected in fixed_rts:
            if rts.get()!=fixed_rts[selected]:rts.set(fixed_rts[selected])
            rts_box.configure(state="disabled")
        else:
            rts_values=["CRS","NIRS","NDRS"] if selected=="Additive" else ["CRS","VRS","NIRS","NDRS"]
            rts_box.configure(state="readonly",values=rts_values)
            if rts.get() not in rts_values:rts.set("CRS")
        if selected in {"SBM","SuperEfficiency","Additive"}:
            if orientation.get()!="Input":orientation.set("Input")
            orientation_box.configure(state="disabled")
        elif selected=="CrossEfficiency":
            if orientation.get()!="Input":orientation.set("Input")
            orientation_box.configure(state="disabled")
        elif selected=="Malmquist_MPI":
            if orientation.get()!="Output":orientation.set("Output")
            orientation_box.configure(state="disabled")
        else:
            orientation_box.configure(state="readonly")
        eq, note = _dea_formula(model.get(), orientation.get(), rts.get())
        if selected=="Additive":note += " Orientation is not used; the normalized slack objective is solved directly."
        if selected=="CrossEfficiency":note += " The multiplier model is CRS and reports aggregate peer scores."
        if selected=="Malmquist_MPI":note += " Uses CRS output distance functions for adjacent periods."
        formula_text.configure(state="normal"); formula_text.delete("1.0","end")
        formula_text.insert("end", f"{model.get()} • {orientation.get()} orientation\n\n{eq}\n\n{note}")
        formula_text.configure(state="disabled")
    for var in (model,orientation,rts): var.trace_add("write",update_formula)
    update_formula()
    ttk.Label(page,text="Solver validation checks DMU IDs, missing values, nonnegative inputs/outputs and model-specific constraints before solving.",style="Muted.TLabel",wraplength=1050).grid(row=4,column=0,sticky="ew",padx=18,pady=4)

    # Results preview and export
    page = pages["results"]; page.columnconfigure(0, weight=1); page.rowconfigure(2, weight=1)
    ttk.Label(page, text="Results Preview", style="Title.TLabel").grid(row=0,column=0,sticky="w",padx=18,pady=(14,2))
    result_info = tk.StringVar(value="Import data, choose a model, then run the analysis.")
    ttk.Label(page,textvariable=result_info,style="Muted.TLabel",wraplength=1080).grid(row=1,column=0,sticky="ew",padx=18,pady=5)
    result_tree = _preview_table(page,height=16,pack=False); result_tree.master.grid(row=2,column=0,sticky="nsew",padx=12,pady=5)
    result_buttons=ttk.Frame(page);result_buttons.grid(row=3,column=0,sticky="ew",padx=18,pady=6)
    def result_frame(result):
        rows=result.get("results",[])
        if rows:
            records=[]
            for r in rows:
                records.append({"DMU_ID":r.dmu,"Period":r.period,"Score":r.score,"Status":r.status,
                                "Input slacks":str(r.input_slacks),"Output slacks":str(r.output_slacks),
                                "Reference peers":str({k:round(v,5) for k,v in r.lambdas.items()})})
            return pd.DataFrame(records)
        if result.get("model")=="Malmquist_MPI":return pd.DataFrame(result.get("malmquist_rows",[]))
        if result.get("model")=="CrossEfficiency":
            matrix=result.get("cross_matrix")
            if matrix is not None:
                df=pd.DataFrame(matrix,index=result.get("names"),columns=result.get("names"));df.insert(0,"Evaluator",df.index);return df.reset_index(drop=True)
        return pd.DataFrame()
    def show_result(result):
        state["result"]=result;df=result_frame(result);_show_df(result_tree,df)
        if df.empty: result_info.set("The model returned no rows. Check the data and model restrictions.")
        else: result_info.set(f"{result['model']} complete • {len(df)} preview rows • scores, peer weights and slacks are shown where applicable. Export includes full result tables and audit notes.")
        app.tabs.select(app.dea_pages["results"])
    def run():
        if not state["bundle"]:
            set_validation("Import a valid workbook in Input & Preview before running DEA.","error");app.error("Import and validate a DEA workbook first.");return
        if model.get()=="SBM" and orientation.get()!="Input": app.error("The supplied SBM solver currently supports Input orientation.");return
        if model.get()=="SuperEfficiency" and orientation.get()!="Input": app.error("The supplied SuperEfficiency solver currently supports Input orientation.");return
        if model.get()=="CrossEfficiency" and rts.get()!="CRS": app.error("The supplied CrossEfficiency solver currently supports CRS.");return
        b=state["bundle"]
        chosen_model,chosen_orientation,chosen_rts=model.get(),orientation.get(),rts.get()
        app.start(lambda:solve_dataset(b["data"],b["input_cols"],b["output_cols"],chosen_model,chosen_orientation,chosen_rts,b["tolerance"]),show_result)
    def export():
        if not state["result"]: app.error("Run the DEA model before exporting results.");return
        path=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel workbook","*.xlsx")])
        if path:
            try: export_results(path,state["result"],state["bundle"]);result_info.set(f"Full result workbook saved: {path}")
            except Exception as exc: app.error(exc)
    app.action(result_buttons,"Run DEA",run,True).pack(side="left",padx=4)
    app.action(result_buttons,"Export full results…",export).pack(side="left",padx=4)
    app.action(result_buttons,"← Input",lambda:app.tabs.select(app.dea_pages["input"])).pack(side="right",padx=4)

    # DEA data generator with preview before saving
    page=pages["generate"];page.columnconfigure(0,weight=1)
    ttk.Label(page,text="Generate DEA Data",style="Title.TLabel").grid(row=0,column=0,sticky="w",padx=18,pady=(14,2))
    ttk.Label(page,text="Create a synthetic DEA workbook for practice. Review the generated rows before saving.",style="Muted.TLabel").grid(row=1,column=0,sticky="w",padx=18,pady=(0,5))
    params=_frame(page,"Generation settings",2);param_vars=[tk.StringVar(value=x) for x in ("30","2","2","42")]
    for i,(label,var) in enumerate(zip(("DMUs (2–2,000)","Inputs (1–10)","Outputs (1–10)","Seed"),param_vars)):
        ttk.Label(params,text=label).grid(row=0,column=i*2,sticky="w",padx=5);ttk.Entry(params,textvariable=var,width=9).grid(row=0,column=i*2+1,sticky="w",padx=5)
    demo_tree=_preview_table(page,height=14,pack=False);demo_tree.master.grid(row=3,column=0,sticky="nsew",padx=12,pady=5);page.rowconfigure(3,weight=1)
    gen_state={"frame":None}
    def generate():
        try:
            df,ins,outs=generate_dea_demo(*(int(v.get()) for v in param_vars))
            gen_state["frame"]=(df,ins,outs);_show_df(demo_tree,df)
            validation.set(f"Generated {len(df)} DMUs with {len(ins)} inputs and {len(outs)} outputs. Values are synthetic and intended for practice.")
        except Exception as exc: app.error(exc)
    def save_generated():
        if not gen_state["frame"]: app.error("Generate data and review the preview first.");return
        path=filedialog.asksaveasfilename(defaultextension=".xlsx",initialfile="DEA_Demo_Input.xlsx",filetypes=[("Excel workbook","*.xlsx")])
        if not path:return
        df,_,_=gen_state["frame"]
        guide=pd.DataFrame({"DEA input template guide":["Replace the yellow input cells on the Data sheet.","DMU_ID: unique identifier for each decision-making unit.","Period: optional for Malmquist analysis.","Input_* columns: resource inputs; use positive numeric values.","Output_* columns: desirable outputs; use nonnegative numeric values.","Generated values are synthetic examples; do not use them as empirical findings."]})
        try:
            with pd.ExcelWriter(path,engine="openpyxl") as writer:
                guide.to_excel(writer,sheet_name="Guide",index=False);df.to_excel(writer,sheet_name="Data",index=False)
                ws=writer.book["Data"]
                from openpyxl.styles import PatternFill
                yellow=PatternFill(fill_type="solid",fgColor="FFF2CC")
                for row in ws.iter_rows(min_row=2,min_col=3):
                    for cell in row:cell.fill=yellow
            import_workbook(path);set_validation(f"Workbook saved and reloaded. Review the imported preview before analysis.")
        except Exception as exc:app.error(exc)
    actions=ttk.Frame(page);actions.grid(row=4,column=0,sticky="ew",padx=18,pady=6)
    app.action(actions,"Generate preview",generate,True).pack(side="left",padx=4)
    app.action(actions,"Save Guide + Data workbook…",save_generated).pack(side="left",padx=4)
    app.action(actions,"← Input",lambda:app.tabs.select(app.dea_pages["input"])).pack(side="right",padx=4)


def build_bootstrap_panels(app, pages):
    """Separate response import from fuzzy scale/formula and keep defuzz output live."""
    state={"input":None,"scale":None,"scale_map":None,"df":None,"xlsx":None,"png":None,"photo":None}
    fields={"B":tk.StringVar(value="2000"),"alpha":tk.StringVar(value="0.05"),"seed":tk.StringVar(value="80"),"start_row":tk.StringVar(value="2"),"start_col":tk.StringVar(value="2"),"header_row":tk.StringVar(value="1"),"n_rows":tk.StringVar(value="0"),"n_cols":tk.StringVar(value="0"),"outname":tk.StringVar(value="bootstrap_result")}
    formula=tk.StringVar(value="(l + m + u) / 3")
    ans_sheet=tk.StringVar(value="Ans 1")

    # Tab 1 — import and inspect expert response data.
    page=pages["input"];page.columnconfigure(0,weight=1);page.rowconfigure(4,weight=1)
    ttk.Label(page,text="Import & Preview Data",style="Title.TLabel").grid(row=0,column=0,sticky="w",padx=18,pady=(10,2))
    ttk.Label(page,text="Load the expert response workbook here. Its Ans sheet and raw values remain visible on this tab.",style="Muted.TLabel",wraplength=1080).grid(row=1,column=0,sticky="w",padx=18,pady=(0,5))
    data_toolbar=ttk.LabelFrame(page,text="Expert response workbook",padding=(9,6));data_toolbar.grid(row=2,column=0,sticky="ew",padx=12,pady=4);data_toolbar.columnconfigure(1,weight=1)
    data_status=tk.StringVar(value="Import a workbook with sheets named Ans 1, Ans 2, …")
    ans_path=tk.StringVar(value="No response workbook selected")
    data_actions=ttk.Frame(data_toolbar);data_actions.grid(row=0,column=0,sticky="w")
    ans_path_label=ttk.Label(data_toolbar,textvariable=ans_path,wraplength=850,style="Muted.TLabel");ans_path_label.grid(row=0,column=1,sticky="ew",padx=8)
    sheetrow=ttk.Frame(page);sheetrow.grid(row=3,column=0,sticky="ew",padx=18,pady=3)
    ttk.Label(sheetrow,text="Preview sheet").pack(side="left",padx=(0,8))
    sheet_pick=ttk.Combobox(sheetrow,textvariable=ans_sheet,state="readonly",width=24);sheet_pick.pack(side="left")
    ans_tree=_preview_table(page,height=19,pack=False);ans_tree.master.grid(row=4,column=0,sticky="nsew",padx=12,pady=4)
    data_status_label=ttk.Label(page,textvariable=data_status,style="Muted.TLabel",wraplength=1080);data_status_label.grid(row=5,column=0,sticky="ew",padx=18,pady=4)
    data_nav=ttk.Frame(page);data_nav.grid(row=6,column=0,sticky="ew",padx=18,pady=(2,8))
    def set_data_status(text,error=False):
        data_status.set(text);data_status_label.configure(style="Validation.Error.TLabel" if error else "Muted.TLabel")
    def refresh_data(*_):
        if not state["input"]:return
        try:
            frame=pd.read_excel(state["input"],sheet_name=ans_sheet.get(),header=None)
            _show_df(ans_tree,frame,250,50)
            set_data_status(f"{ans_sheet.get()} preview • {frame.shape[0]} rows × {frame.shape[1]} columns.")
        except Exception as exc:
            _show_df(ans_tree,pd.DataFrame());set_data_status(f"Could not preview this sheet: {exc}",True)
    sheet_pick.bind("<<ComboboxSelected>>",refresh_data)

    # Tab 2 — fuzzy scale, formula, and the result of applying defuzzification.
    page=pages["fuzzy"];page.columnconfigure(0,weight=2);page.columnconfigure(1,weight=3);page.rowconfigure(4,weight=1)
    ttk.Label(page,text="Import & Preview Fuzzy Scale",style="Title.TLabel").grid(row=0,column=0,columnspan=2,sticky="w",padx=18,pady=(10,2))
    ttk.Label(page,text="Import the scale, configure the defuzzification formula, then inspect the crisp matrix generated from the selected Ans sheet.",style="Muted.TLabel",wraplength=1080).grid(row=1,column=0,columnspan=2,sticky="w",padx=18,pady=(0,5))
    scale_toolbar=ttk.LabelFrame(page,text="Fuzzy scale workbook",padding=(9,6));scale_toolbar.grid(row=2,column=0,columnspan=2,sticky="ew",padx=12,pady=4);scale_toolbar.columnconfigure(1,weight=1)
    scale_status=tk.StringVar(value="No fuzzy scale selected")
    scale_path=tk.StringVar(value="Import the workbook that maps each label to L, M and U.")
    scale_actions=ttk.Frame(scale_toolbar);scale_actions.grid(row=0,column=0,sticky="w")
    ttk.Label(scale_toolbar,textvariable=scale_path,wraplength=850,style="Muted.TLabel").grid(row=0,column=1,sticky="ew",padx=8)
    scalebox=_frame(page,"Fuzzy scale preview (Label, L, M, U)",3,0)
    formula_box=_frame(page,"Defuzzification configuration",3,1)
    scalebox.columnconfigure(0,weight=1);scalebox.rowconfigure(0,weight=1)
    formula_box.columnconfigure(0,weight=1)
    scale_tree=_preview_table(scalebox,height=11,pack=False);scale_tree.master.grid(row=0,column=0,sticky="nsew",padx=5,pady=4)
    ttk.Label(formula_box,text="Formula for each (l, m, u) triple",font=(app.ui_font,10,"bold")).grid(row=0,column=0,sticky="w",padx=6,pady=(3,2))
    formula_entry=ttk.Entry(formula_box,textvariable=formula);formula_entry.grid(row=1,column=0,sticky="ew",padx=6,pady=3)
    formula_help="Allowed: l, m, u; numbers; +, -, *, /, ** and parentheses. Example: (l + 2*m + u) / 4."
    ttk.Label(formula_box,text=formula_help,wraplength=560,justify="left",style="Muted.TLabel").grid(row=2,column=0,sticky="ew",padx=6,pady=3)
    formula_status=tk.StringVar(value="Import both workbooks to see defuzzified values.")
    formula_status_label=ttk.Label(formula_box,textvariable=formula_status,style="Muted.TLabel",wraplength=560)
    formula_status_label.grid(row=3,column=0,sticky="ew",padx=6,pady=3)
    ttk.Label(formula_box,textvariable=scale_status,style="Muted.TLabel",wraplength=560).grid(row=4,column=0,sticky="ew",padx=6,pady=3)
    defuzz_box=_frame(page,"Defuzzified preview — selected Ans sheet",4,0,colspan=2)
    defuzz_box.columnconfigure(0,weight=1);defuzz_box.rowconfigure(0,weight=1)
    defuzz_tree=_preview_table(defuzz_box,height=14,pack=False);defuzz_tree.master.grid(row=0,column=0,sticky="nsew",padx=5,pady=4)
    fuzzy_nav=ttk.Frame(page);fuzzy_nav.grid(row=5,column=0,columnspan=2,sticky="ew",padx=18,pady=(2,8))

    def refresh_defuzz_preview(*_):
        if not state["input"] or not state["scale_map"]:
            _show_df(defuzz_tree,pd.DataFrame())
            formula_status.set("Import the response workbook and fuzzy scale to calculate the preview.")
            formula_status_label.configure(style="Muted.TLabel");return
        try:
            fn=build_defuzzify_func(formula.get())
            raw=pd.read_excel(state["input"],sheet_name=ans_sheet.get(),header=None)
            sr=int(fields["start_row"].get());sc=int(fields["start_col"].get());hr=int(fields["header_row"].get())
            if min(sr,sc,hr)<1:raise ValueError("Excel row and column settings start at 1.")
            header=list(raw.iloc[hr-1,sc-1:].tolist())
            labels=[]
            for value in header:
                if pd.isna(value) or str(value).strip()=="":break
                labels.append(str(value).strip())
            nc=int(fields["n_cols"].get()) or len(labels);nr=int(fields["n_rows"].get()) or len(labels)
            if nr<1 or nc<1:raise ValueError("No factor headers found. Check the header row and data start column settings.")
            labels=labels[:nc]
            if len(labels)<nc:raise ValueError(f"Expected {nc} factor headers but found only {len(labels)}.")
            matrix=[];row_names=[]
            for i in range(nr):
                excel_row=sr-1+i
                if excel_row>=len(raw):raise ValueError("The configured data range extends past the selected Ans sheet.")
                row_label=raw.iat[excel_row,sc-2] if sc>1 and sc-2<raw.shape[1] else None
                row_names.append(str(row_label).strip() if pd.notna(row_label) and str(row_label).strip() else (labels[i] if i<len(labels) else f"Row {i+1}"))
                values=[]
                for j in range(nc):
                    excel_col=sc-1+j
                    if excel_col>=raw.shape[1]:raise ValueError("The configured data range extends past the selected Ans sheet.")
                    value=raw.iat[excel_row,excel_col]
                    triple=parse_fuzzy_label(value,state["scale_map"])
                    values.append(fn(*triple))
                matrix.append(values)
            result=pd.DataFrame(matrix,columns=labels);result.insert(0,"Factor",row_names);_show_df(defuzz_tree,result,250,50)
            formula_status.set(f"Formula applied: {nr} × {nc} crisp values from {ans_sheet.get()} • {formula.get()}.")
            formula_status_label.configure(style="Validation.Ok.TLabel")
        except Exception as exc:
            _show_df(defuzz_tree,pd.DataFrame());formula_status.set(f"Could not calculate preview: {exc}")
            formula_status_label.configure(style="Validation.Error.TLabel")
    formula.trace_add("write",refresh_defuzz_preview)
    for key in ("start_row","start_col","header_row","n_rows","n_cols"):
        fields[key].trace_add("write",refresh_defuzz_preview)

    def validate_response_labels():
        if not state["input"] or not state["scale_map"]:return
        try:
            fn=build_defuzzify_func(formula.get())
            nr=int(fields["n_rows"].get()) or None;nc=int(fields["n_cols"].get()) or None
            _,matrices=read_expert_sheets(state["input"],start_row=int(fields["start_row"].get()),start_col=int(fields["start_col"].get()),n_rows=nr,n_cols=nc,header_row=int(fields["header_row"].get()),fuzzy_scale=state["scale_map"],defuzz_func=fn)
            scale_status.set(f"Label validation passed • {len(matrices)} experts • {matrices[0].shape[0]} factors.")
            refresh_defuzz_preview()
        except Exception as exc:scale_status.set(f"Response labels need attention: {exc}")

    def import_ans():
        path=filedialog.askopenfilename(filetypes=[("Excel workbooks","*.xlsx *.xlsm")])
        if not path:return
        try:
            book=pd.ExcelFile(path);sheets=[s for s in book.sheet_names if s.strip().lower().startswith("ans ")]
            if not sheets:raise ValueError("No Ans 1, Ans 2, … expert sheets found. Check the workbook sheet names.")
            state["input"]=path;ans_path.set(f"{Path(path).name} • {len(sheets)} expert sheets")
            sheet_pick.configure(values=sheets);ans_sheet.set(sheets[0]);refresh_data();refresh_defuzz_preview();validate_response_labels()
        except Exception as exc:
            set_data_status(f"Workbook validation failed: {exc}",True);app.error(f"Bootstrap response workbook validation failed:\n\n{exc}")
    def import_scale():
        path=filedialog.askopenfilename(filetypes=[("Excel workbooks","*.xlsx *.xlsm")])
        if not path:return
        try:
            mapping=load_fuzzy_scale(path)
            if not mapping:raise ValueError("No fuzzy scale rows found. Expected columns A–D: label, L, M, U.")
            state.update(scale=path,scale_map=mapping);scale_path.set(f"{Path(path).name} • {len(mapping)} fuzzy labels")
            scale_df=pd.DataFrame([{"Label":k,"L":v[0],"M":v[1],"U":v[2]} for k,v in mapping.items()]);_show_df(scale_tree,scale_df,200,10)
            scale_status.set(f"Loaded {len(mapping)} fuzzy labels.");refresh_defuzz_preview();validate_response_labels()
        except Exception as exc:
            scale_status.set(f"Fuzzy scale validation failed: {exc}");app.error(f"Fuzzy scale validation failed:\n\n{exc}")
    app.action(data_actions,"Import response workbook…",import_ans,True).pack(side="left",padx=(0,6))
    app.action(scale_actions,"Import fuzzy scale…",import_scale,True).pack(side="left",padx=(0,6))
    app.action(data_nav,"Next: fuzzy scale & formula →",lambda:app.tabs.select(pages["fuzzy"]),True).pack(side="right")
    app.action(fuzzy_nav,"← Response data",lambda:app.tabs.select(pages["input"])).pack(side="left")
    app.action(fuzzy_nav,"Continue to Bootstrap settings →",lambda:app.tabs.select(pages["settings"]),True).pack(side="right")

    # Tab 3 — bootstrap run configuration (scale and formula live on tab 2).
    page=pages["settings"];page.columnconfigure(0,weight=1)
    ttk.Label(page,text="Bootstrap Configuration",style="Title.TLabel").grid(row=0,column=0,sticky="w",padx=18,pady=(10,2))
    ttk.Label(page,text="Set the resampling design and the response data range. Fuzzy mapping and defuzzification are configured on the preceding tab.",style="Muted.TLabel",wraplength=1080).grid(row=1,column=0,sticky="w",padx=18,pady=(0,5))
    settings=_frame(page,"Bootstrap and data range",2);settings.columnconfigure(1,weight=1);settings.columnconfigure(3,weight=1);settings.columnconfigure(5,weight=1)
    rows=[("Bootstrap samples B","B"),("Confidence alpha","alpha"),("Random seed","seed"),("Data start row","start_row"),("Data start column","start_col"),("Factor header row","header_row"),("Number of factors (0 = auto)","n_rows"),("Number of columns (0 = auto)","n_cols"),("Output base name","outname")]
    for i,(label,key) in enumerate(rows):
        r=i//3;c=(i%3)*2
        ttk.Label(settings,text=label).grid(row=r,column=c,sticky="w",padx=5,pady=5)
        ttk.Entry(settings,textvariable=fields[key],width=18).grid(row=r,column=c+1,sticky="ew",padx=5,pady=5)
    settings_status=tk.StringVar(value="Import both workbooks, review the crisp preview, then run Bootstrap.")
    ttk.Label(page,textvariable=settings_status,style="Muted.TLabel",wraplength=1050).grid(row=3,column=0,sticky="ew",padx=18,pady=6)
    actions=ttk.Frame(page);actions.grid(row=4,column=0,sticky="ew",padx=18,pady=6)
    def start_bootstrap():
        if not state["input"] or not state["scale_map"]:app.error("Import the response workbook and fuzzy scale first.");return
        try:
            b=int(fields["B"].get());a=float(fields["alpha"].get());seed=int(fields["seed"].get())
            sr=int(fields["start_row"].get());sc=int(fields["start_col"].get());hr=int(fields["header_row"].get())
            nr=int(fields["n_rows"].get()) or None;nc=int(fields["n_cols"].get()) or None
            if b<1:raise ValueError("Bootstrap samples B must be at least 1.")
            if not 0<a<1:raise ValueError("Confidence alpha must be between 0 and 1.")
            if min(sr,sc,hr)<1 or (nr is not None and nr<1) or (nc is not None and nc<1):raise ValueError("Excel rows and columns start at 1; factor counts must be positive or 0 for auto-detect.")
            fn=build_defuzzify_func(formula.get());outdir=Path(tempfile.mkdtemp(prefix="datastudio_bootstrap_"));stem=Path(fields["outname"].get().strip() or "bootstrap_result").name;stem=Path(stem).stem or "bootstrap_result"
            state["xlsx"]=outdir/f"{stem}.xlsx";state["png"]=outdir/f"{stem}.png"
            args=dict(B=b,alpha=a,seed=seed,start_row=sr,start_col=sc,n_rows=nr,n_cols=nc,header_row=hr,fuzzy_scale=dict(state["scale_map"]),defuzz_func=fn)
            settings_status.set("Bootstrap is running…")
            app.start(lambda:run_pipeline(state["input"],state["xlsx"],state["png"],**args),show_boot_result)
        except Exception as exc:app.error(exc)
    app.action(actions,"Run Bootstrap",start_bootstrap,True).pack(side="left",padx=4)
    app.action(actions,"← Fuzzy scale & formula",lambda:app.tabs.select(pages["fuzzy"])).pack(side="right",padx=4)

    # Tab 4 — analysis output and exports.
    page=pages["results"];page.columnconfigure(0,weight=1);page.rowconfigure(2,weight=1)
    ttk.Label(page,text="Bootstrap Results",style="Title.TLabel").grid(row=0,column=0,sticky="w",padx=18,pady=(10,2))
    result_status=tk.StringVar(value="Run Bootstrap from Bootstrap Configuration to preview the result table and IRM figure.")
    ttk.Label(page,textvariable=result_status,style="Muted.TLabel",wraplength=1080).grid(row=1,column=0,sticky="ew",padx=18,pady=5)
    split=ttk.Panedwindow(page,orient="horizontal");split.grid(row=2,column=0,sticky="nsew",padx=12,pady=5)
    table_frame=ttk.Frame(split);image_frame=ttk.LabelFrame(split,text="IRM visualization",padding=8);split.add(table_frame,weight=3);split.add(image_frame,weight=2)
    result_tree=_preview_table(table_frame,height=15)
    image_label=ttk.Label(image_frame,text="IRM preview appears after analysis.");image_label.pack(fill="both",expand=True)
    buttons=ttk.Frame(page);buttons.grid(row=3,column=0,sticky="ew",padx=18,pady=7)
    def show_boot_result(df):
        state["df"]=df;_show_df(result_tree,df,250,40)
        try:
            from PIL import Image,ImageTk
            img=Image.open(state["png"]);img.thumbnail((700,480));state["photo"]=ImageTk.PhotoImage(img)
            image_label.configure(image=state["photo"],text="")
        except Exception as exc:image_label.configure(text=f"The analysis completed but the plot preview could not be loaded: {exc}")
        result_status.set(f"Bootstrap complete • {len(df)} factors • B={fields['B'].get()} • alpha={fields['alpha'].get()}. Export the workbook and IRM image separately.")
        app.tabs.select(pages["results"])
    def export_file(kind):
        src=state["xlsx"] if kind=="Excel" else state["png"]
        if not src or not src.exists():app.error("Run Bootstrap before exporting files.");return
        ext=".xlsx" if kind=="Excel" else ".png";types=[("Excel workbook","*.xlsx")] if kind=="Excel" else [("PNG image","*.png")]
        dest=filedialog.asksaveasfilename(defaultextension=ext,filetypes=types,initialfile=src.name)
        if dest:
            try:shutil.copy2(src,dest);result_status.set(f"{kind} exported to {dest}")
            except Exception as exc:app.error(exc)
    app.action(buttons,"Save Excel result…",lambda:export_file("Excel"),True).pack(side="left",padx=4)
    app.action(buttons,"Save IRM image…",lambda:export_file("Image")).pack(side="left",padx=4)
    app.action(buttons,"← Bootstrap settings",lambda:app.tabs.select(pages["settings"])).pack(side="right",padx=4)
