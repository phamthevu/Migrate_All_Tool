import os,json
import re
import tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def export_excel(path, dfs, cancel=None):
    if not dfs: raise ValueError('No data to export.')
    names=set()
    wb=Workbook(write_only=False); wb.remove(wb.active)
    for name,df in dfs.items():
        if cancel and cancel.is_set():raise InterruptedError('Cancelled; no output file written.')
        if not name or len(name)>31 or re.search(r'[\\/*?:\[\]]',name) or name.startswith("'") or name.endswith("'") or name.casefold() in names:
            raise ValueError(f'Invalid or duplicate sheet name: {name}')
        names.add(name.casefold())
        ws=wb.create_sheet(name)
        for values in [[df.index.name or 'Row']+list(df.columns)]+[[idx]+row.tolist() for idx,row in df.iterrows()]:
            ws.append([json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict,tuple)) else v for v in values])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
        ws.freeze_panes='B2'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='245A81')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(c.value or '')) for c in col)+2))
    dest=Path(path)
    fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=dest.parent); os.close(fd)
    try:
        wb.save(tmp)
        if cancel and cancel.is_set():raise InterruptedError('Cancelled; no output file written.')
        os.replace(tmp,dest)
    finally:
        if os.path.exists(tmp): os.unlink(tmp)
