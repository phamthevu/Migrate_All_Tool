from __future__ import annotations
from pathlib import Path
import pandas as pd
import numpy as np
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

APP_DIR=Path(__file__).resolve().parents[1]
ROOT=APP_DIR.parent
def shared_root():
    return Path(sys.executable).resolve().parent if getattr(sys,'frozen',False) else ROOT
TEMPLATE_DIR=shared_root()/'Templates'/'Template'/'MCDMStudio'
DEMO_DIR=shared_root()/'Templates'/'Demo Data'/'MCDMStudio'
TOPSIS_TEMPLATE=TEMPLATE_DIR/'MCDM_TOPSIS_Input_Template.xlsx'
DEMATEL_TEMPLATE=TEMPLATE_DIR/'MCDM_DEMATEL_Input_Template.xlsx'
TEMPLATE=TOPSIS_TEMPLATE
DEMO=DEMO_DIR/'MCDM_Demo_Data.xlsx'
DEMATEL_DEMO=DEMO_DIR/'MCDM_DEMATEL_Demo_Data.xlsx'

SHEET_METHODS={
    'Instructions':('All methods','Workbook layout, input rules and source notes.'),
    'Method Guide':('All methods','Quick lookup of which worksheets each method needs.'),
    'Decision Matrix':('SAW, TOPSIS, MOORA, MULTIMOORA, COPRAS-G, ARAS, WASPAS, VIKOR, ELECTRE I, PROMETHEE II, CoCoSo, EDAS, MARCOS, CODAS, MAIRCA, MABAC, GRA, PIV; Entropy, CRITIC, MEREC, Standard deviation, PSI weights','Alternative rows and criterion columns; also used by data-based objective weighting methods.'),
    'Criteria':('All ranking methods; AHP, ANP, SWARA, BWM, FUCOM, Entropy, CRITIC, MEREC, Standard deviation, PSI weights; ISM','Criterion names; Benefit/Cost direction and manual weights for matrix methods; BWM comparisons where marked.'),
    'AHP Pairwise':('AHP','Pairwise criterion comparisons.'),
    'SWARA':('SWARA','Criterion priority order and relative importance values.'),
    'FUCOM':('FUCOM','Criterion order and comparative priorities between adjacent criteria.'),
    'DEMATEL Direct Relations':('DEMATEL','Square factor-by-factor direct influence matrix (0–4; diagonal 0).'),
    'ISM SSIM':('ISM','Pairwise contextual relation judgments using V/A/X/O.'),
    'ANP Weighted Supermatrix':('ANP','Weighted, column-stochastic supermatrix.'),
}
METHOD_INPUT_SHEETS={
    'AHP':'AHP Pairwise','ANP':'ANP Weighted Supermatrix','SWARA':'SWARA','BWM':'Criteria','FUCOM':'FUCOM',
    'DEMATEL':'DEMATEL Direct Relations','ISM':'ISM SSIM',
    'Entropy weights':'Decision Matrix','CRITIC':'Decision Matrix','MEREC':'Decision Matrix',
    'Standard deviation':'Decision Matrix','PSI weights':'Decision Matrix',
}

def method_input_preview(data, method):
    """Return the actual input worksheet for the first displayed calculation step."""
    sheet=METHOD_INPUT_SHEETS.get(method,'Decision Matrix')
    previews=data.get('sheet_previews',{})
    frame=previews.get(sheet)
    if frame is not None and not frame.empty:return frame.copy()
    return data.get('preview',pd.DataFrame()).copy()

def _add_method_guide(wb, kind):
    guide=wb.create_sheet('Method Guide')
    guide.append(['Sheet','Used by method(s)','Purpose / expected data'])
    sheets=['Instructions','Method Guide','DEMATEL Direct Relations'] if kind=='DEMATEL' else ['Instructions','Method Guide','Decision Matrix','Criteria','AHP Pairwise','SWARA','FUCOM','ANP Weighted Supermatrix','DEMATEL Direct Relations','ISM SSIM']
    for sheet in sheets:
        methods,purpose=SHEET_METHODS[sheet]
        guide.append([sheet,methods,purpose])
    _style(guide,input_from_col=99)
    guide.column_dimensions['A'].width=30; guide.column_dimensions['B'].width=90; guide.column_dimensions['C'].width=76
    for row in guide.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    for ws in wb.worksheets:
        methods,purpose=SHEET_METHODS.get(ws.title,('See Method Guide',''))
        ws['A1'].comment=Comment(f'Methods: {methods}\nPurpose: {purpose}\nSee the Method Guide sheet for the complete mapping.','MCDMStudio')

SOURCE_URLS={
    'AHP':'https://doi.org/10.1007/978-981-33-4745-8','ANP':'https://doi.org/10.1007/978-981-33-4745-8','TOPSIS':'https://doi.org/10.1007/978-981-33-4745-8',
    'DEMATEL':'https://doi.org/10.1007/978-981-33-4745-8','ISM':'https://doi.org/10.1109/TSMC.1974.4309336','FUCOM':'https://doi.org/10.3390/sym10090393',
    'BWM':'https://doi.org/10.1016/j.omega.2014.11.009','CRITIC':'https://doi.org/10.1016/0305-0548(94)00059-H','MEREC':'https://doi.org/10.3390/sym13040525',
    'PSI weights':'https://doi.org/10.1016/j.matdes.2009.11.020','CoCoSo':'https://doi.org/10.1108/MD-05-2017-0458','EDAS':'https://doi.org/10.15388/Informatica.2015.57',
    'MARCOS':'https://doi.org/10.1016/j.cie.2019.106231','MABAC':'https://doi.org/10.1016/j.eswa.2014.11.057','PIV':'https://doi.org/10.1016/j.cie.2018.03.045',
}

def _style(ws, header=1, input_from_row=2, input_from_col=2):
    ws.freeze_panes=f'B{header+1}'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[header]:
        cell.font=Font(color='FFFFFF',bold=True,size=10); cell.fill=PatternFill('solid',fgColor='173B57'); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[header].height=30
    for row in ws.iter_rows(min_row=input_from_row):
        for cell in row:
            if cell.column>=input_from_col:
                cell.fill=PatternFill('solid',fgColor='FFF2C6'); cell.font=Font(color='173B57')
                cell.alignment=Alignment(vertical='center',horizontal='center')
            elif cell.value is not None:
                cell.fill=PatternFill('solid',fgColor='EAF1F7'); cell.font=Font(bold=True,color='29465B')
    for col in ws.columns:
        letter=col[0].column_letter; width=min(30,max(12,max((len(str(c.value or '')) for c in col),default=0)+2)); ws.column_dimensions[letter].width=width

def generate_template(path=TEMPLATE, demo=False, template_type='TOPSIS'):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); kind=str(template_type).upper()
    if kind not in {'TOPSIS','DEMATEL'}: raise ValueError("Template type must be 'TOPSIS' or 'DEMATEL'.")
    wb=Workbook(); info=wb.active; info.title='Instructions'
    if kind=='DEMATEL':
        rows=[['MCDMStudio | DEMATEL factor influence template','Input cells are pale yellow. Rows and columns must list the same factors in the same order.'],['Purpose','Record direct influence from one factor (row) to another factor (column).'],['Scale','Use 0–4: no influence (0) to very high influence (4).'],['Diagonal','Set factor-to-itself influence to 0. Diagonal cells are shaded gray.'],['Analysis','Import this workbook, review and edit values in MCDMStudio, save, then continue to DEMATEL.'],['Source','Thakkar (2021), Multi-Criteria Decision Making, Chapter 9, §§9.2–9.3.']]
        for row in rows: info.append(row)
        matrix=wb.create_sheet('DEMATEL Direct Relations'); factors=['F1','F2','F3','F4']; matrix.append(['Factor',*factors])
        values=[[0,2,1,0],[1,0,3,1],[0,2,0,2],[1,0,1,0]] if demo else [[0,'','',''],['',0,'',''],['','',0,''],['','','',0]]
        for name,row in zip(factors,values): matrix.append([name,*row])
        info['A1'].font=Font(size=15,bold=True,color='173B57')
        info.column_dimensions['A'].width=27; info.column_dimensions['B'].width=100
        for cell in info[1]: cell.fill=PatternFill('solid',fgColor='173B57'); cell.font=Font(size=12 if cell.column==1 else 10,bold=True,color='FFFFFF'); cell.alignment=Alignment(wrap_text=True,vertical='center')
        info.row_dimensions[1].height=36
        for r in range(2,info.max_row+1):
            info.cell(r,1).fill=PatternFill('solid',fgColor='EAF1F7'); info.cell(r,1).font=Font(bold=True,color='29465B')
            info.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center'); info.row_dimensions[r].height=28
        _add_method_guide(wb,kind)
        _style(matrix)
        for i in range(2,matrix.max_row+1):
            diagonal=i
            c=matrix.cell(i,diagonal); c.fill=PatternFill('solid',fgColor='D8E1E8'); c.font=Font(color='718096',italic=True)
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines=False
            for col in ws.columns:
                letter=col[0].column_letter; ws.column_dimensions[letter].width=max(ws.column_dimensions[letter].width or 0,min(32,max(12,max((len(str(c.value or '')) for c in col),default=0)+3)))
        wb.save(path); return path
    rows=[['MCDMStudio | TOPSIS-style decision matrix','Input cells are pale yellow. Rows are alternatives; columns are criteria/factors.'],['Purpose','Enter numeric performance values for each alternative and factor.'],['Step 1 - Decision Matrix','One alternative per row; one numeric factor per column.'],['Step 2 - Criteria','Use Benefit when higher is better and Cost when lower is better.'],['Step 3 - Weighting','Choose weights in the app or enter manual weights in Criteria.'],['Step 4 - Save','Choose a method, edit any relevant sheet in the app, then save and calculate in one action.'],['Sources','TOPSIS: Thakkar (2021), Chapter 5. AHP: Chapter 3. See Method Guide and the app for sheet mappings and references.']]
    for row in rows: info.append(row)
    info.column_dimensions['A'].width=28; info.column_dimensions['B'].width=105
    for r in range(2,info.max_row+1):
        info.cell(r,1).fill=PatternFill('solid',fgColor='EAF1F7'); info.cell(r,1).font=Font(bold=True,color='29465B')
        info.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center'); info.row_dimensions[r].height=28
    info['A1'].font=Font(size=15,bold=True,color='173B57')
    for cell in info[1]: cell.fill=PatternFill('solid',fgColor='173B57'); cell.font=Font(size=12 if cell.column==1 else 10,bold=True,color='FFFFFF'); cell.alignment=Alignment(wrap_text=True,vertical='center')
    info.row_dimensions[1].height=36
    matrix=wb.create_sheet('Decision Matrix'); matrix.append(['Alternative','F1','F2','F3','F4'])
    if demo:
        for row in [['A1',80,70,75,60],['A2',65,90,80,70],['A3',90,60,65,85],['A4',75,80,90,75]]: matrix.append(row)
    else:
        for i in range(1,5): matrix.append([f'Alternative {i}','','','',''])
    crit=wb.create_sheet('Criteria'); crit.append(['Criterion','Direction','Manual weight','BWM best?','BWM worst?','Best-to-criterion','Criterion-to-worst']);
    bwm=[('Y','',1,7),('','Y',7,1),('','',3,5),('','',5,3)] if demo else [('', '',1,1) for _ in range(4)]
    for (c,d,w),bw in zip([('F1','Benefit',0.3),('F2','Cost',0.25),('F3','Benefit',0.25),('F4','Benefit',0.2)],bwm): crit.append([c,d,w,*bw])
    pair=wb.create_sheet('AHP Pairwise'); pair.append(['Criterion','F1','F2','F3','F4'])
    vals=[[1,1,1,1],[1,1,1,1],[1,1,1,1],[1,1,1,1]]
    if demo: vals=[[1,3,2,4],[1/3,1,1/2,2],[1/2,2,1,3],[1/4,1/2,1/3,1]]
    for c,row in zip(['F1','F2','F3','F4'],vals): pair.append([c]+row)
    sw=wb.create_sheet('SWARA'); sw.append(['Rank','Criterion','s (relative importance)'])
    for i,c in enumerate(['F1','F2','F3','F4'],1): sw.append([i,c,0 if i==1 else 0.1])
    fu=wb.create_sheet('FUCOM'); fu.append(['Rank','Criterion','Comparative priority to next'])
    for i,c in enumerate(['F1','F2','F3','F4'],1): fu.append([i,c,1.0 if i==4 else (1.5 if i==1 else 1.2)])
    direct=wb.create_sheet('DEMATEL Direct Relations'); direct.append(['Factor','F1','F2','F3','F4'])
    vals=[[0,2,1,0],[1,0,3,1],[0,2,0,2],[1,0,1,0]] if demo else [[0,'','',''],['',0,'',''],['','',0,''],['','','',0]]
    for c,row in zip(['F1','F2','F3','F4'],vals): direct.append([c]+row)
    ssim=wb.create_sheet('ISM SSIM'); ssim.append(['Element','C1','C2','C3','C4'])
    vals=[['','V','V','V'],['','', 'V','V'],['','','','V'],['','','','']] if demo else [['','','',''],['','','',''],['','','',''],['','','','']]
    for c,row in zip(['F1','F2','F3','F4'],vals): ssim.append([c]+row)
    anp=wb.create_sheet('ANP Weighted Supermatrix'); anp.append(['Element','C1','C2','C3','C4'])
    vals=[[.4,.2,.2,.1],[.3,.4,.2,.2],[.2,.2,.4,.3],[.1,.2,.2,.4]] if demo else [[.25,.25,.25,.25] for _ in range(4)]
    for c,row in zip(['F1','F2','F3','F4'],vals): anp.append([c]+row)
    _add_method_guide(wb,kind)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines=False
        if ws.title not in ('Instructions','Method Guide'):
            _style(ws)
            if ws.title=='DEMATEL Direct Relations':
                for i in range(2,ws.max_row+1):
                    c=ws.cell(i,i); c.fill=PatternFill('solid',fgColor='D8E1E8'); c.font=Font(color='718096',italic=True)
    wb.save(path); return path

def save_edited_input(path, output_path, preview, decision_mode='TOPSIS'):
    path=Path(path); output_path=Path(output_path)
    wb=load_workbook(path)
    sheet='DEMATEL Direct Relations' if decision_mode=='DEMATEL' else 'Decision Matrix'
    if sheet not in wb.sheetnames: raise ValueError(f"Workbook is missing '{sheet}'.")
    ws=wb[sheet]
    df=preview.copy()
    # Keep labels/headers stable, update only editable performance cells.
    for ri in range(len(df)):
        for ci in range(1,len(df.columns)):
            value=df.iloc[ri,ci]
            try: value=float(value)
            except (TypeError,ValueError): raise ValueError(f'Cell {ri+2},{ci+1} must contain a number.')
            if not np.isfinite(value): raise ValueError('Values must be finite numbers.')
            ws.cell(row=ri+2,column=ci+1,value=value)
    output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return output_path

def save_workbook_edits(path, output_path, sheet_frames):
    """Write edits from the in-app sheet editor into a reviewed workbook copy."""
    path=Path(path); output_path=Path(output_path); wb=load_workbook(path)
    for sheet,frame in sheet_frames.items():
        if sheet not in wb.sheetnames: raise ValueError(f"Workbook is missing '{sheet}'.")
        ws=wb[sheet]
        for col_idx, header in enumerate(frame.columns, start=1):
            ws.cell(row=1,column=col_idx,value=str(header))
        for row_idx,row in enumerate(frame.itertuples(index=False,name=None),start=2):
            for col_idx,value in enumerate(row,start=1):
                if pd.isna(value) or (isinstance(value,str) and not value.strip()): value=None
                elif isinstance(value,str):
                    try:
                        number=float(value)
                        value=int(number) if number.is_integer() else number
                    except ValueError: pass
                ws.cell(row=row_idx,column=col_idx,value=value)
    output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return output_path

def load_input(path):
    xls=pd.ExcelFile(path)
    # Keep every worksheet available to the UI's workbook preview. Calculation
    # inputs are still parsed below from their named method sheets.
    sheet_previews={name:pd.read_excel(xls,sheet_name=name).dropna(how='all') for name in xls.sheet_names}
    # Dedicated DEMATEL workbooks use factors on both axes and need no alternatives sheet.
    if 'Decision Matrix' not in xls.sheet_names and 'DEMATEL Direct Relations' in xls.sheet_names:
        raw=pd.read_excel(path,sheet_name='DEMATEL Direct Relations').dropna(how='all')
        if raw.shape[0]<2 or raw.shape[1]<3: raise ValueError('DEMATEL workbook needs at least two factors.')
        factors=raw.iloc[:,0].astype(str).tolist(); headers=[str(v) for v in raw.columns[1:]]
        if len(factors)!=len(headers) or factors!=headers: raise ValueError('DEMATEL factor labels on rows and columns must match exactly and use the same order.')
        raw_values=raw.iloc[:,1:].apply(pd.to_numeric,errors='coerce')
        bad=raw.iloc[:,1:].notna() & raw_values.isna()
        if bad.any().any(): raise ValueError('DEMATEL matrix contains non-numeric values.')
        matrix=raw_values.to_numpy(dtype=float)
        preview=raw.copy(); preview.columns=['Factor',*headers]
        result={'path':str(path),'decision_mode':'DEMATEL','alternatives':factors,'criteria':factors,'directions':['Benefit']*len(factors),'matrix':matrix,'weights':[1/len(factors)]*len(factors),'swara':{},'ahp':None,'bwm':None,'dematel':matrix,'ism':None,'anp':None,'fucom':None,'preview':preview,'sheet_previews':sheet_previews}
        xls.close(); return result
    if 'Decision Matrix' not in xls.sheet_names or 'Criteria' not in xls.sheet_names: raise ValueError("Workbook must include 'Decision Matrix' and 'Criteria' sheets, or a square 'DEMATEL Direct Relations' sheet.")
    dm=pd.read_excel(path,sheet_name='Decision Matrix').dropna(how='all')
    meta=pd.read_excel(path,sheet_name='Criteria').dropna(how='all')
    if dm.shape[1]<3 or meta.shape[1]<2: raise ValueError('The workbook needs an Alternative column and at least two criteria.')
    alt_col=dm.columns[0]; alternatives=dm[alt_col].astype(str).tolist()
    criteria=[str(c) for c in dm.columns[1:]]
    if meta.shape[0]!=len(criteria): raise ValueError('Criteria sheet must have one row for each decision-matrix criterion.')
    # map by criterion name while accepting order-aligned template labels
    names_col=meta.columns[0]; direction_col=meta.columns[1]
    rowmap={str(row[names_col]):row for _,row in meta.iterrows()}
    if set(criteria)<=set(rowmap): rows=[rowmap[c] for c in criteria]
    else: rows=[row for _,row in meta.iterrows()]
    directions=[str(row[direction_col]).strip().title() for row in rows]
    raw_matrix=dm.iloc[:,1:]
    matrix=raw_matrix.apply(pd.to_numeric,errors='coerce')
    bad=raw_matrix.notna() & matrix.isna()
    if bad.any().any(): raise ValueError('Decision Matrix contains non-numeric performance values.')
    if len(set(alternatives))!=len(alternatives): raise ValueError('Alternative names must be unique.')
    weight_col=meta.columns[2] if meta.shape[1]>=3 else None
    weights=[]
    for row in rows:
        v=row[weight_col] if weight_col else None
        weights.append(float(v) if pd.notna(v) else 0.)
    if sum(weights)<=0: weights=[1/len(criteria)]*len(criteria)
    swara={}
    if 'SWARA' in xls.sheet_names:
        sdf=pd.read_excel(path,sheet_name='SWARA').dropna(how='all')
        if sdf.shape[1]>=3:
            sdf=sdf.sort_values(sdf.columns[0]); swara={str(r.iloc[1]):float(r.iloc[2]) for _,r in sdf.iterrows() if pd.notna(r.iloc[1]) and pd.notna(r.iloc[2])}
    fucom=None
    if 'FUCOM' in xls.sheet_names:
        fdf=pd.read_excel(path,sheet_name='FUCOM').dropna(how='all')
        if fdf.shape[1]>=3:
            fdf=fdf.dropna(subset=[fdf.columns[1]]).sort_values(fdf.columns[0])
            fucom={'order':[str(v) for v in fdf.iloc[:,1]],'priorities':[pd.to_numeric(v,errors='coerce') for v in fdf.iloc[:-1,2]]}
    ahp=None
    if 'AHP Pairwise' in xls.sheet_names:
        pdf=pd.read_excel(path,sheet_name='AHP Pairwise')
        if pdf.shape[0]>=len(criteria) and pdf.shape[1]>=len(criteria)+1:
            ahp=pdf.iloc[:len(criteria),1:len(criteria)+1].apply(pd.to_numeric,errors='coerce').to_numpy()
    def read_square(sheet, numeric=False):
        if sheet not in xls.sheet_names:return None
        frame=pd.read_excel(path,sheet_name=sheet).dropna(how='all')
        if frame.shape[0]<len(criteria) or frame.shape[1]<len(criteria)+1:return None
        raw=frame.iloc[:len(criteria),1:len(criteria)+1]
        if numeric:
            vals=raw.apply(pd.to_numeric,errors='coerce').to_numpy(dtype=float)
            return vals if np.isfinite(vals).all() else None
        return raw.fillna('').to_numpy(dtype=object)
    bwm=None
    required={'BWM best?','BWM worst?','Best-to-criterion','Criterion-to-worst'}
    if required.issubset(set(meta.columns)):
        by_name={str(row[names_col]):row for _,row in meta.iterrows()}
        best_names=[str(row[names_col]) for _,row in meta.iterrows() if str(row['BWM best?']).strip().upper() in ('Y','YES','1','TRUE')]
        worst_names=[str(row[names_col]) for _,row in meta.iterrows() if str(row['BWM worst?']).strip().upper() in ('Y','YES','1','TRUE')]
        if len(best_names)==1 and len(worst_names)==1 and all(c in by_name for c in criteria):
            bwm={'best_idx':criteria.index(best_names[0]),'worst_idx':criteria.index(worst_names[0]),'best_to_others':[float(by_name[c]['Best-to-criterion']) for c in criteria],'others_to_worst':[float(by_name[c]['Criterion-to-worst']) for c in criteria]}
    result={'path':str(path),'decision_mode':'TOPSIS','alternatives':alternatives,'criteria':criteria,'directions':directions,'matrix':matrix.to_numpy(dtype=float),'weights':weights,'swara':swara,'fucom':fucom,'ahp':ahp,'bwm':bwm,'dematel':read_square('DEMATEL Direct Relations',True),'ism':read_square('ISM SSIM',False),'anp':read_square('ANP Weighted Supermatrix',True),'preview':dm,'sheet_previews':sheet_previews}
    xls.close()
    return result

def write_results(path, method, input_data, weights, steps, params=None):
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        source_url=SOURCE_URLS.get(method,'')
        summary=pd.DataFrame({'Property':['Method','Method source URL','Input file','Alternative / factor count','Criterion / factor count','Weight source','Parameters'],'Value':[method,source_url,input_data.get('path',''),len(input_data['alternatives']),len(input_data['criteria']),input_data.get('weight_source','Manual / template'),str(params or {})]})
        summary.to_excel(writer,sheet_name='Summary',index=False)
        if source_url: writer.book['Summary']['B3'].hyperlink=source_url; writer.book['Summary']['B3'].style='Hyperlink'
        pd.DataFrame({'Criterion':input_data['criteria'],'Direction':input_data['directions'],'Weight':weights}).to_excel(writer,sheet_name='Weights',index=False)
        pd.DataFrame(input_data['matrix'],index=input_data['alternatives'],columns=input_data['criteria']).rename_axis('Alternative').to_excel(writer,sheet_name='Input Matrix')
        for i,step in enumerate(steps,1):
            df=step.table.copy()
            if df.index.name is None and not isinstance(df.index,pd.RangeIndex): df=df.reset_index()
            sheet=f'Step {i:02d}'
            df.to_excel(writer,sheet_name=sheet,index=False)
            ws=writer.book[sheet]; ws.cell(row=1,column=max(1,len(df.columns)+1),value='Formula').font=__import__('openpyxl').styles.Font(bold=True,color='174C69')
            ws.cell(row=2,column=max(1,len(df.columns)+1),value=step.formula)
            ws.cell(row=3,column=max(1,len(df.columns)+1),value='Book reference: '+step.book_ref)
            ws.cell(row=4,column=max(1,len(df.columns)+1),value='Source URL: '+source_url if source_url else 'Source URL: see README references')
            if source_url:
                link_cell=ws.cell(row=4,column=max(1,len(df.columns)+1)); link_cell.hyperlink=source_url; link_cell.style='Hyperlink'
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=__import__('openpyxl').styles.Font(color='FFFFFF',bold=True); cell.fill=__import__('openpyxl').styles.PatternFill('solid',fgColor='174C69')
            for col in ws.columns:
                letter=col[0].column_letter; ws.column_dimensions[letter].width=min(45,max(12,max((len(str(c.value or '')) for c in col),default=0)+2))
