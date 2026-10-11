import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from MCDMStudio.core.excel_io import DEMO, load_input, write_results
from MCDMStudio.core.methods import calculate_ahp, calculate_swara, run_ranking

class MCDMMethodTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load_input(DEMO)

    def test_demo_workbook_loads(self):
        self.assertEqual(self.data['matrix'].shape,(4,4))
        self.assertEqual(len(self.data['criteria']),4)
        self.assertEqual(set(self.data['directions']),{'Benefit','Cost'})

    def test_ahp_weights_are_normalized_and_consistent(self):
        weights, steps=calculate_ahp(self.data['criteria'],self.data['ahp'])
        self.assertAlmostEqual(float(weights.sum()),1.0,places=10)
        self.assertLess(float(steps[-1].table.iloc[0]['CR']),.1)

    def test_ahp_rejects_nonreciprocal_matrix(self):
        matrix=self.data['ahp'].copy(); matrix[0,1]=5
        with self.assertRaisesRegex(ValueError,'reciprocal'):
            calculate_ahp(self.data['criteria'],matrix)

    def test_swara_weights_are_normalized(self):
        weights, steps=calculate_swara(self.data['criteria'],self.data['swara'])
        self.assertAlmostEqual(float(weights.sum()),1.0,places=10)
        self.assertIn('Weight',steps[-1].table.columns)

    def test_cost_direction_in_saw(self):
        steps=run_ranking('SAW',[[10,1],[10,10]],['low cost','high cost'],['Benefit','Cost'],['Benefit','Cost'],[.5,.5])
        self.assertEqual(steps[-1].table.iloc[0]['Alternative'],'low cost')

    def test_every_ranking_method_returns_stepwise_tables(self):
        for method in ['SAW','TOPSIS','MOORA','COPRAS-G','ARAS','WASPAS','VIKOR']:
            with self.subTest(method=method):
                steps=run_ranking(method,self.data['matrix'],self.data['alternatives'],self.data['criteria'],self.data['directions'],self.data['weights'])
                self.assertGreaterEqual(len(steps),3)
                self.assertTrue(all(step.title and step.formula and step.book_ref for step in steps))
                self.assertGreaterEqual(len(steps[-1].table),2)
        moora=run_ranking('MOORA',self.data['matrix'],self.data['alternatives'],self.data['criteria'],self.data['directions'],self.data['weights'])[-1].table
        self.assertIn('Alternative',moora.columns)
        self.assertIn('Rank',moora.columns)

    def test_export_creates_one_sheet_per_step(self):
        steps=run_ranking('TOPSIS',self.data['matrix'],self.data['alternatives'],self.data['criteria'],self.data['directions'],self.data['weights'])
        with tempfile.TemporaryDirectory() as temp:
            out=Path(temp)/'steps.xlsx'; write_results(out,'TOPSIS',self.data,self.data['weights'],steps)
            sheets=load_workbook(out,read_only=True).sheetnames
            self.assertEqual(len(sheets),len(steps)+3)
            self.assertIn('Summary',sheets); self.assertIn('Input Matrix',sheets)

if __name__=='__main__': unittest.main()
