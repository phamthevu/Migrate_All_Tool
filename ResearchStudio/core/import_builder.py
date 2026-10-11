from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment

COLORS={'required':'F4CCCC','data':'CFE2F3','optional':'FFF2CC','guide':'D9EAD3','system':'D9D9D9'}

def _s(v): return '' if v is None else str(v).strip()
def _bool(v): return str(v).strip().lower() in {'1','true','yes','y','on'}
def _style(ws, kind):
    ws.sheet_properties.tabColor=COLORS[kind]
    for c in ws[1]: c.fill=PatternFill('solid',fgColor=COLORS[kind]);c.font=Font(bold=True);c.alignment=Alignment(vertical='center')
    ws.freeze_panes='A2'
def _finish(wb,path):
    if 'START_HERE' in wb.sheetnames: wb._sheets.insert(0,wb._sheets.pop(wb._sheets.index(wb['START_HERE'])))
    for ws in wb.worksheets:
        for col in ws.columns:
            try: ws.column_dimensions[col[0].column_letter].width=min(42,max(12,max(len(str(c.value or '')) for c in col)+2))
            except Exception: pass
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)

def validate_definition(d):
    e=[]; pt=d.get('project_type','')
    if pt not in {'MCDM','QUANTITATIVE_SURVEY','QUALITATIVE_INTERVIEW'}: e.append('Choose a supported project type.')
    if pt=='MCDM':
        if not d.get('factors'): e.append('Add at least one Factor.')
        if d.get('model')=='TOPSIS' and not d.get('alternatives'): e.append('TOPSIS requires at least one Solution / Alternative.')
        if len(d.get('scale',[]))<2: e.append('Add at least two scale rows.')
        if d.get('use_z') and len(d.get('z_scale',[]))<2: e.append('Z-number requires reliability/certainty scale rows.')
    elif pt=='QUANTITATIVE_SURVEY':
        qs=d.get('questions',[])
        if not qs:e.append('Add at least one quantitative column/question.')
        ids=[_s(x.get('question_id')) for x in qs];
        if any(not x for x in ids):e.append('Every quantitative row needs Question_ID.')
        if len(ids)!=len(set(ids)):e.append('Question_ID must be unique.')
        parents=set(ids)
        matrix_items=d.get('matrix_items',[])
        choice_types={'SINGLE_CHOICE','MULTI_CHOICE','LIKERT','MATRIX_LIKERT'}
        for q in qs:
            qid=_s(q.get('question_id'));qtype=_s(q.get('question_type')).upper();opts=_s(q.get('options'))
            if qtype in choice_types and not opts:
                e.append(f"{qid}: {qtype} requires Options / Choices separated by |.")
            if qtype=='MATRIX_LIKERT' and not any(_s(m.get('parent_id'))==qid for m in matrix_items):
                e.append(f"{qid}: MATRIX_LIKERT requires at least one matrix component.")
        for m in matrix_items:
            if _s(m.get('parent_id')) not in parents:e.append(f"Matrix parent '{_s(m.get('parent_id'))}' is not a configured question.")
            if not _s(m.get('item_id')):e.append('Every matrix component needs an Item_ID.')
            if not _s(m.get('item_text')):e.append(f"Matrix component '{_s(m.get('item_id'))}' needs a label/statement.")
    elif pt=='QUALITATIVE_INTERVIEW':
        qs=d.get('questions',[])
        if not qs:e.append('Add at least one interview question.')
        ids=[_s(x.get('question_id')) for x in qs]
        if any(not x for x in ids):e.append('Every interview row needs Question_ID.')
        if len(ids)!=len(set(ids)):e.append('Question_ID must be unique.')
    return e

def summary_definition(d):
    pt=d.get('project_type'); mode=d.get('source_mode','UI')
    if pt=='MCDM': return f"READY — MCDM {d.get('model','DEMATEL')} | Z={'ON' if d.get('use_z') else 'OFF'} | {len(d.get('factors',[]))} factors | {len(d.get('alternatives',[]))} solutions | {len(d.get('scale',[]))} scale rows | source={mode}"
    if pt=='QUANTITATIVE_SURVEY': return f"READY — Quantitative | {len(d.get('questions',[]))} questions/columns | {len(d.get('matrix_items',[]))} matrix components | source={mode}"
    return f"READY — Qualitative | {len(d.get('questions',[]))} initial questions | source={mode}"

def infer_from_sample(path, project_type, model='DEMATEL', use_z=False):
    wb=load_workbook(path,data_only=True,read_only=True);pt=project_type
    d={'project_type':pt,'model':model,'use_z':bool(use_z),'source_mode':'SAMPLE','sample_file':Path(path).name,'access_mode':'Free'}
    if pt=='MCDM':
        ans=[s for s in wb.sheetnames if str(s).upper().startswith('ANS')]
        if not ans: raise ValueError('MCDM sample requires at least one sheet whose name starts with ANS (for example Ans_01).')
        rows=list(wb[ans[0]].iter_rows(values_only=True,max_row=100))
        if len(rows)<2: raise ValueError('The ANS sheet must contain a header and at least one matrix row.')
        cols=[_s(x) for x in rows[0]][1:]; rids=[_s(r[0]) for r in rows[1:] if _s(r[0])]
        factors=[];alts=[]
        if model=='TOPSIS':
            factors=[{'code':x,'name':x} for x in cols if x];alts=[{'code':x,'name':x} for x in rids]
        else:
            seen=[]
            for x in rids+cols:
                if x and x not in seen:seen.append(x)
            factors=[{'code':x,'name':x} for x in seen]
        vals=[]
        for r in rows[1:]:
            for v in r[1:]:
                x=_s(v)
                if x and x not in vals:vals.append(x)
        acodes=[];bcodes=[]
        for x in vals:
            a,b=(x.split('|',1)+[''])[:2] if '|' in x else (x,'')
            if a and a not in acodes:acodes.append(a)
            if b and b not in bcodes:bcodes.append(b)
        d.update(factors=factors,alternatives=alts,scale=[{'term':x,'code':x,'value':''} for x in acodes] or [{'term':'Low','code':'L','value':''},{'term':'High','code':'H','value':''}],z_scale=[{'term':x,'code':x,'value':''} for x in bcodes],sample_ans=ans[0])
    elif pt=='QUANTITATIVE_SURVEY':
        sn='Data' if 'Data' in wb.sheetnames else wb.sheetnames[0];rows=list(wb[sn].iter_rows(values_only=True,max_row=25))
        if len(rows)<2: raise ValueError('Quantitative sample requires a header row and at least one data row.')
        hdr=[_s(x) for x in rows[0]];ignore={'ANSWER_ID','RESPONDENT_ID','PARTICIPANT_NAME','USERNAME','ORGANIZATION','SUBMITTED_AT','MODIFIED_AT'}
        qs=[]
        for j,h in enumerate(hdr):
            if not h or h.upper() in ignore:continue
            vals=[r[j] for r in rows[1:] if j<len(r) and r[j] not in (None,'')]
            typ='SHORT_TEXT'
            if vals and all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in vals):typ='INTEGER' if all(float(v).is_integer() for v in vals) else 'DECIMAL'
            qs.append({'question_id':h,'question_text':h.replace('_',' ').title(),'question_type':typ,'required':False,'options':''})
        if not qs: raise ValueError('No usable response columns were detected in the sample header.')
        d.update(questions=qs,matrix_items=[],sample_sheet=sn)
    else:
        sn='Data' if 'Data' in wb.sheetnames else wb.sheetnames[0];rows=list(wb[sn].iter_rows(values_only=True,max_row=200))
        if len(rows)<2:raise ValueError('Qualitative sample requires a header row and at least one answer row.')
        hdr=[_s(x) for x in rows[0]]; up={x.upper():i for i,x in enumerate(hdr)}
        if 'ANSWER' not in up:raise ValueError("Qualitative sample must contain an 'Answer' column.")
        qs=[]
        if 'QUESTION_ID' in up:
            qi=up['QUESTION_ID']; qt=up.get('QUESTION');seen=set()
            for r in rows[1:]:
                qid=_s(r[qi] if qi<len(r) else '')
                if qid and qid not in seen:
                    seen.add(qid);qs.append({'question_id':qid,'question_text':_s(r[qt]) if qt is not None and qt<len(r) else qid,'required':False})
        else:
            # Wide legacy table: every non-identity column becomes an initial question.
            ignore={'RESPONDENT_ID','PARTICIPANT_NAME','USERNAME','ORGANIZATION','SUBMITTED_AT','MODIFIED_AT','ANSWER'}
            for h in hdr:
                if h and h.upper() not in ignore:qs.append({'question_id':h,'question_text':h.replace('_',' ').title(),'required':False})
        if not qs:qs=[{'question_id':'Q01','question_text':'Imported interview question','required':False}]
        d.update(questions=qs,sample_sheet=sn)
    errors=validate_definition(d)
    if errors: raise ValueError('\n'.join(errors))
    return d

def generate_data_template(path,d):
    errors=validate_definition(d)
    if errors:raise ValueError('\n'.join(errors))
    wb=Workbook();wb.remove(wb.active);pt=d['project_type']
    g=wb.create_sheet('START_HERE');g.append(['RESEARCH STUDIO V8.5 — DATA ENTRY TEMPLATE']);g.append(['Step','Instruction']);
    for x in [('1','Configuration/schema has already been prepared in Research Studio.'),('2','Paste historical data ONLY into BLUE data tabs. Do not rename headers or tabs.'),('3','Example rows are illustrative; replace/delete them before import.'),('4','Return to Import Excel → Preview & Create Project.'),('Color','BLUE=data entry; GREEN=reference; YELLOW=optional identity/account; RED=required reference.')]:g.append(x)
    _style(g,'guide')
    sys=wb.create_sheet('_System');sys.append(['Key','Value']);
    for k,v in [('Template_Version','8.5'),('Project_Type',pt),('Model_Type',d.get('model','')),('Use_Z_Number','YES' if d.get('use_z') else 'NO'),('Access_Mode',d.get('access_mode','Free'))]:sys.append([k,v])
    _style(sys,'system');sys.sheet_state='hidden'
    ps=wb.create_sheet('Project_Setup');ps.append(['Setting','Value']);ps.append(['Project Name',d.get('project_name','Imported Historical Project')]);ps.append(['Access Mode',d.get('access_mode','Free')]);_style(ps,'guide')
    if pt=='QUANTITATIVE_SURVEY':
        q=wb.create_sheet('Questions');q.append(['Question_ID','Question_Text','Question_Type','Required','Options','Min_Value','Max_Value','Default_Value','Validation','Description','Display_Order'])
        for i,x in enumerate(d['questions'],1):q.append([x.get('question_id'),x.get('question_text'),x.get('question_type','SHORT_TEXT'),bool(x.get('required')),x.get('options',''),x.get('min_value',''),x.get('max_value',''),'','','',i])
        _style(q,'guide');mi=wb.create_sheet('Matrix_Items');mi.append(['Parent_ID','Item_ID','Item_Text','Display_Order'])
        for i,x in enumerate(d.get('matrix_items',[]),1):mi.append([x.get('parent_id'),x.get('item_id'),x.get('item_text'),i])
        _style(mi,'guide');ac=wb.create_sheet('Accounts');ac.append(['Respondent_ID','Username','Password','Full_Name','Organization']);_style(ac,'optional')
        ids=[]
        mids={x.get('parent_id') for x in d.get('matrix_items',[])}
        for qx in d['questions']:
            qid=qx['question_id'];children=[x['item_id'] for x in d.get('matrix_items',[]) if x.get('parent_id')==qid]
            ids.extend(children or [qid])
        da=wb.create_sheet('Data');da.append(['Answer_ID','Respondent_ID','Participant_Name','Username','Organization','Submitted_At','Modified_At']+ids);da.append(['Ans_001','R001','Example participant','','','','']+['' for _ in ids]);_style(da,'data')
    elif pt=='QUALITATIVE_INTERVIEW':
        q=wb.create_sheet('Questions');q.append(['Question_ID','Question_Text','Required','Display_Order'])
        for i,x in enumerate(d['questions'],1):q.append([x['question_id'],x.get('question_text',''),bool(x.get('required')),i])
        _style(q,'guide');ac=wb.create_sheet('Accounts');ac.append(['Respondent_ID','Username','Password','Full_Name','Organization']);_style(ac,'optional')
        da=wb.create_sheet('Data');da.append(['Interview_ID','Respondent_ID','Participant_Name','Username','Organization','Submitted_At','Modified_At','Sequence','Type','Question_ID','Question','Answer']);q0=d['questions'][0];da.append(['Interview_001','R001','Example participant','','','','',1,'INITIAL',q0['question_id'],q0.get('question_text',''),'']);_style(da,'data')
    else:
        fa=wb.create_sheet('Factors_Alternatives');fa.append(['Type','Code','Name','Description'])
        for x in d['factors']:fa.append(['Factor',x['code'],x.get('name',x['code']),''])
        for x in d.get('alternatives',[]):fa.append(['Alternative',x['code'],x.get('name',x['code']),''])
        _style(fa,'guide');ss=wb.create_sheet('Scale_Standard');ss.append(['Linguistic_Term','Code','Value'])
        for x in d['scale']:ss.append([x.get('term'),x.get('code'),x.get('value','')])
        _style(ss,'guide');sz=wb.create_sheet('Scale_Z_Number');sz.append(['A_Linguistic_Term','A_Code','A_Value','B_Linguistic_Term','B_Code','B_Value'])
        if d.get('use_z'):
            for i,x in enumerate(d.get('z_scale',[])):sz.append(['','','',x.get('term'),x.get('code'),x.get('value','')])
        _style(sz,'guide');ex=wb.create_sheet('Experts');ex.append(['Respondent_ID','Expert_Name','Username','Password','Organization']);ex.append(['E001','Example expert','expert01','123','']);_style(ex,'optional')
        lk=wb.create_sheet('Lookup');lk.append(['Answer_ID','Respondent_ID','Expert_Name','Username','Organization','Submitted_At','Modified_At']);lk.append(['Ans_01','E001','Example expert','expert01','','','']);_style(lk,'data')
        rows=[x['code'] for x in (d.get('alternatives') if d.get('model')=='TOPSIS' else d['factors'])];cols=[x['code'] for x in d['factors']]
        if d.get('model')=='DELPHI':cols=['Rating']
        ans=wb.create_sheet('Ans_01');ans.append(['']+cols)
        for r in rows:ans.append([r]+['' for _ in cols])
        _style(ans,'data')
    sc=wb.create_sheet('Schema');sc.append(['Source','Project_Type','Item','Meaning'])
    sc.append([d.get('source_mode','UI'),pt,'Configuration','Prepared/validated by Research Studio V8.5']);_style(sc,'guide')
    # comments on all data headers
    for ws in wb.worksheets:
        if ws.sheet_properties.tabColor and str(ws.sheet_properties.tabColor.rgb or '').endswith(COLORS['data']):
            for c in ws[1]: c.comment=Comment('Keep this header unchanged. Paste data below it.','Research Studio')
    return _finish(wb,path)
