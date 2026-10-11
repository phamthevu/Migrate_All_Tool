from __future__ import annotations
import json
from pathlib import Path
from openpyxl import Workbook
from core.models import config_from_dict
from core.excel_handler import export_all_experts_results

LAYOUT={'TITLE','SECTION','DESCRIPTION'}

def project_type(db): return str((db.get_meta() or {}).get('project_type') or 'MCDM').upper()
def _safe(v): return '' if v is None else v

def _legacy_answer(v):
    # V7.4 no longer creates quantitative follow-ups, but old projects remain readable.
    if isinstance(v,dict) and ('answer' in v or 'followups' in v): return v.get('answer')
    return v

def _quant_ids(config):
    ids=[]
    for q in config.get('questions',[]):
        if str(q.get('question_type','')).upper() in LAYOUT: continue
        if str(q.get('question_type','')).upper().startswith('MATRIX_'):
            ids.extend(str(x.get('item_id')) for x in q.get('matrix_items',[]) if x.get('item_id'))
        elif q.get('question_id'): ids.append(str(q.get('question_id')))
    return ids

def _quant_value(payload,var_id,questions):
    if var_id in payload:
        v=_legacy_answer(payload.get(var_id)); return json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v
    for q in questions:
        if str(q.get('question_type','')).upper().startswith('MATRIX_'):
            nested=_legacy_answer(payload.get(q.get('question_id'))) or {}
            if isinstance(nested,dict) and var_id in nested:return nested.get(var_id)
    return None

def export_quantitative(db,path):
    cfg=db.get_config_dict(); qs=cfg.get('questions',[]); ids=_quant_ids(cfg); wb=Workbook(); ws=wb.active;ws.title='Data'
    ws.append(['Answer_ID','Respondent_ID','Participant_Name','Username','Organization','Source_Project_ID','Source_Project','Source_File','Submitted_At','Modified_At']+ids)
    for r in db.get_all_responses():
        d=json.loads(r['response_json'])
        ws.append([f"Ans_{int(r['display_index']):02d}",r.get('respondent_id'),r.get('full_name'),r.get('username'),r.get('organization'),r.get('source_project_id'),r.get('source_project_name'),r.get('source_file'),r.get('submitted_at'),r.get('modified_at')]+[_quant_value(d,i,qs) for i in ids])
    code=wb.create_sheet('Codebook');code.append(['Variable_ID','Question_Text','Question_Type','Required','Options','Parent_ID'])
    for q in qs:
        if str(q.get('question_type','')).upper() in LAYOUT:continue
        code.append([q.get('question_id'),q.get('question_text'),q.get('question_type'),q.get('required'),' | '.join(map(str,q.get('options',[]) or [])),''])
        for mi in q.get('matrix_items',[]) or []:code.append([mi.get('item_id'),mi.get('item_text'),'MATRIX_ITEM',q.get('required'),' | '.join(map(str,q.get('options',[]) or [])),q.get('question_id')])
    schema=wb.create_sheet('Schema'); schema.append(['Field_Name','Role','Data_Type','Required','Description']);
    for x in ['Respondent_ID','Participant_Name','Username','Organization','Source_Project_ID','Source_Project','Source_File','Submitted_At','Modified_At']: schema.append([x,'METADATA','TEXT',False,'Response metadata']);
    for x in ids: schema.append([x,'DATA','TEXT',False,'Configured survey variable / matrix item']);
    resp=wb.create_sheet('Respondents');resp.append(['Respondent_ID','Full_Name','Username','Organization','Status'])
    for r in db.get_respondents():resp.append([r.get('respondent_id'),r.get('full_name'),r.get('username'),r.get('organization'),r.get('status')])
    meta=wb.create_sheet('Metadata');
    for k,v in (db.get_meta() or {}).items():
        if k!='config_json':meta.append([k,v])
    for sh in wb.worksheets:sh.freeze_panes='A2'
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)

def export_qualitative(db,path):
    cfg=db.get_config_dict();wb=Workbook();ws=wb.active;ws.title='Data';ws.append(['Interview_ID','Respondent_ID','Participant_Name','Username','Organization','Source_Project_ID','Source_Project','Source_File','Submitted_At','Modified_At','Sequence','Type','Question_ID','Question','Answer'])
    for r in db.get_all_responses():
        iid=f"Interview_{int(r['display_index']):02d}";items=json.loads(r['response_json']).get('interview_items',[])
        if not items:ws.append([iid,r.get('respondent_id'),r.get('full_name'),r.get('username'),r.get('organization'),r.get('source_project_id'),r.get('source_project_name'),r.get('source_file'),r.get('submitted_at'),r.get('modified_at'),'','','','',''])
        for x in items:ws.append([iid,r.get('respondent_id'),r.get('full_name'),r.get('username'),r.get('organization'),r.get('source_project_id'),r.get('source_project_name'),r.get('source_file'),r.get('submitted_at'),r.get('modified_at'),x.get('sequence'),x.get('kind'),x.get('question_id'),x.get('question'),x.get('answer')])
    schema=wb.create_sheet('Schema');schema.append(['Field_Name','Role','Data_Type','Required','Description']);
    for x in ['Respondent_ID','Participant_Name','Username','Organization','Source_Project_ID','Source_Project','Source_File','Submitted_At','Modified_At','Sequence','Type','Question_ID','Question','Answer']:schema.append([x,'TRANSCRIPT' if x in {'Sequence','Type','Question_ID','Question','Answer'} else 'METADATA','TEXT',False,'Qualitative export/import vocabulary']);
    guide=wb.create_sheet('Initial_Guide');guide.append(['Question_ID','Question_Text','Required','Display_Order'])
    for q in cfg.get('initial_questions',[]):guide.append([q.get('question_id'),q.get('question_text'),q.get('required'),q.get('display_order')])
    meta=wb.create_sheet('Metadata')
    for k,v in (db.get_meta() or {}).items():
        if k!='config_json':meta.append([k,v])
    for sh in wb.worksheets:sh.freeze_panes='A2'
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);return str(path)

def export_mcdm(db,raw_path,num_path):
    cfg=config_from_dict(db.get_config_dict());experts={}
    for idx,r in enumerate(db.get_all_responses(),1):
        p=json.loads(r['response_json']);rid=str(r.get('respondent_id') or r.get('id'))
        experts[rid]={'answer_id':f"Ans_{int(r.get('display_index') or idx):02d}",'respondent_id':rid,'expert_name':r.get('full_name',''),'username':r.get('username',''),'source_project_id':r.get('source_project_id',''),'source_project_name':r.get('source_project_name',''),'submitted_at':r.get('submitted_at',''),'modified_at':r.get('modified_at',''),'source_file':r.get('source_file',''),'code_matrix':p.get('code_matrix',[]),'value_matrix':p.get('value_matrix',[])}
    f=[x.code for x in cfg.factors];a=[x.code for x in cfg.alternatives];m=cfg.model_type.upper()
    rows,cols=(f,f) if m in {'DEMATEL','AHP'} else ((a,f) if m=='TOPSIS' else ((f,['Rating']) if m=='DELPHI' else ([],[])))
    export_all_experts_results(str(raw_path),str(num_path),cfg,experts,rows,cols);return [str(raw_path),str(num_path)]
