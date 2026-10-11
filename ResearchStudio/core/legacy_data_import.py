from pathlib import Path
from datetime import datetime
import json, uuid
from openpyxl import Workbook, load_workbook

META=['Respondent_ID','Username','Full_Name','Organization']
SCHEMA_COLS=['Field_Name','Role','Data_Type','Required','Source_Column','Description']
ALIASES={'Participant_Name':'Full_Name','Expert_Name':'Full_Name','RespondentID':'Respondent_ID','Participant_ID':'Respondent_ID'}

def _norm(v): return ''.join(ch.lower() for ch in _s(v) if ch.isalnum())
def _add_schema_sheet(wb, fields, project_type):
    ws=wb.create_sheet('Schema'); ws.append(SCHEMA_COLS)
    for name,role,dtype,req,desc in fields:
        ws.append([name,role,dtype,bool(req),name,desc])
    ws['H1']='Project_Type'; ws['H2']=project_type
    ws['I1']='Usage'; ws['I2']='Source_Column may be edited to map a historical column to Field_Name when Auto-detect is enabled.'

def _schema_mapping(path, expected):
    wb=load_workbook(path,data_only=True,read_only=True)
    mapping={x:x for x in expected}
    if 'Schema' not in wb.sheetnames: return mapping
    rows=list(wb['Schema'].iter_rows(values_only=True))
    if not rows:return mapping
    h=[_s(x) for x in rows[0]]; idx={x:i for i,x in enumerate(h)}
    for row in rows[1:]:
        field=_s(row[idx.get('Field_Name',0)]) if row else ''
        src=_s(row[idx.get('Source_Column',4)]) if len(row)>idx.get('Source_Column',4) else ''
        if field in mapping and src: mapping[field]=src
    return mapping

def _remap_records(headers, rows, expected, path, autodetect):
    if not autodetect:return headers,rows,{}
    mapping=_schema_mapping(path,expected)
    norm_to_header={_norm(h):h for h in headers}
    used={}
    for field in expected:
        source=mapping.get(field,field)
        actual=source if source in headers else norm_to_header.get(_norm(source))
        if not actual:
            # safe built-in aliases only
            for alias,target in ALIASES.items():
                if target==field and _norm(alias) in norm_to_header:
                    actual=norm_to_header[_norm(alias)]; break
        if actual: used[field]=actual
    remapped=[]
    for r in rows:
        nr=dict(r)
        for field,actual in used.items(): nr[field]=r.get(actual)
        remapped.append(nr)
    return list(dict.fromkeys(headers+list(used))),remapped,used

def _s(v): return '' if v is None else str(v).strip()
def _unique_nonblank(values, label):
    xs=[x for x in values if x]
    if len(xs)!=len(set(xs)): raise ValueError(f'Duplicate {label} values are not allowed.')

def _save(wb,path):
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(max(12,max(len(str(c.value or '')) for c in col)+2),45)
    Path(path).parent.mkdir(parents=True,exist_ok=True); wb.save(path)

def generate_import_template(db,path):
    meta=db.get_meta() or {}; cfg=db.get_config_dict(); pt=meta.get('project_type')
    wb=Workbook(); ws=wb.active; ws.title='Data'
    if pt=='QUANTITATIVE_SURVEY':
        cols=[]
        for q in cfg.get('questions',[]):
            if str(q.get('question_type','')).startswith('MATRIX_'):
                cols += [str(x.get('item_id')) for x in q.get('matrix_items',[]) if x.get('item_id')]
            elif q.get('question_id') and str(q.get('question_type','')) not in {'TITLE','SECTION','DESCRIPTION'}:
                cols.append(str(q['question_id']))
        ws.append(META+cols)
        ws.append(['R001','legacy001','Legacy Participant 01','']+['' for _ in cols])
    elif pt=='QUALITATIVE_INTERVIEW':
        ws.append(META+['Sequence','Kind','Parent_Question_ID','Question_ID','Question','Answer'])
        q=(cfg.get('initial_questions') or [{}])[0]
        ws.append(['R001','legacy001','Legacy Participant 01','', '1','INITIAL','',q.get('question_id','Q01'),q.get('question_text',''),''])
        ws.append(['R001','legacy001','Legacy Participant 01','', '1.1','FOLLOW_UP',q.get('question_id','Q01'),'','Follow-up question',''])
    else:
        rows=[]; cols=[]
        factors=[x for x in cfg.get('factors',[]) if str(x.get('type','Factor')).lower()!='alternative']
        alts=[x for x in cfg.get('alternatives',[])]
        model=str(cfg.get('model_type') or cfg.get('model') or '').upper()
        if model=='TOPSIS': rows=alts; cols=factors
        elif model=='DELPHI': rows=factors; cols=[{'code':'Rating'}]
        else: rows=factors; cols=factors
        z=bool(cfg.get('use_z_number'))
        ws.append(META+['Row_Code','Column_Code','A_Code']+(['B_Code'] if z else []))
        rc=_s((rows[0] if rows else {}).get('code')) or 'C1'; cc=_s((cols[0] if cols else {}).get('code')) or 'C1'
        ws.append(['R001','legacy001','Legacy Expert 01','',rc,cc,'']+([''] if z else []))
    # Schema is shared vocabulary for import/preview/export. It does not modify project config.
    fields=[(x,'IDENTITY','TEXT',x=='Respondent_ID','Participant / respondent metadata') for x in META]
    if pt=='QUANTITATIVE_SURVEY':
        fields += [(x,'DATA','TEXT',False,'Configured survey variable / matrix item') for x in cols]
    elif pt=='QUALITATIVE_INTERVIEW':
        fields += [(x,'TRANSCRIPT','TEXT',x in {'Sequence','Kind','Answer'},'Qualitative transcript field') for x in ['Sequence','Kind','Parent_Question_ID','Question_ID','Question','Answer']]
    else:
        fields += [(x,'MCDM','TEXT',x in {'Row_Code','Column_Code','A_Code'},'MCDM matrix field') for x in ['Row_Code','Column_Code','A_Code']+(['B_Code'] if bool(cfg.get('use_z_number')) else [])]
    _add_schema_sheet(wb,fields,pt)
    info=wb.create_sheet('README'); info.append(['Research Studio V8.1 Historical Data Import Template']); info.append(['Project',meta.get('project_name','')]); info.append(['Project Type',pt]); info.append(['Storage Key',meta.get('storage_key','')]); info.append(['Rule','Do not rename headers. One Respondent_ID identifies one response. Blank answer cells are allowed unless your research requires them.'])
    _save(wb,path)

def _records(path):
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        if 'Data' not in wb.sheetnames: raise ValueError("Missing required sheet 'Data'. Download a fresh template from this project.")
        ws=wb['Data']; rows=list(ws.iter_rows(values_only=True))
        if not rows: raise ValueError('Data sheet is empty.')
        headers=[_s(x) for x in rows[0]]
        if len(headers)!=len(set(headers)): raise ValueError('Duplicate column headers detected.')
        return headers,[dict(zip(headers,r)) for r in rows[1:] if any(_s(x) for x in r)]
    finally:
        wb.close()

def validate_import_file(db,path,autodetect=False):
    headers,rows=_records(path); meta=db.get_meta() or {}; cfg=db.get_config_dict(); pt=meta.get('project_type')
    expected=list(META)
    if pt=='QUANTITATIVE_SURVEY':
        for q in cfg.get('questions',[]):
            if str(q.get('question_type','')).startswith('MATRIX_'): expected += [str(x.get('item_id')) for x in q.get('matrix_items',[]) if x.get('item_id')]
            elif q.get('question_id') and str(q.get('question_type','')) not in {'TITLE','SECTION','DESCRIPTION'}: expected.append(str(q['question_id']))
    elif pt=='QUALITATIVE_INTERVIEW': expected += ['Sequence','Kind','Parent_Question_ID','Question_ID','Question','Answer']
    else: expected += ['Row_Code','Column_Code','A_Code']+(['B_Code'] if bool(cfg.get('use_z_number')) else [])
    headers,rows,mapping=_remap_records(headers,rows,expected,path,autodetect)
    errors=[]; warnings=[]; preview=[]
    if autodetect:
        warnings.append('Auto-detect schema enabled. Mapped columns: '+(', '.join(f'{k} <- {v}' for k,v in mapping.items()) or 'none'))
    for c in META:
        if c not in headers: errors.append(f'Missing required column: {c}')
    if pt=='QUANTITATIVE_SURVEY':
        expected=[]
        for q in cfg.get('questions',[]):
            if str(q.get('question_type','')).startswith('MATRIX_'): expected += [str(x.get('item_id')) for x in q.get('matrix_items',[]) if x.get('item_id')]
            elif q.get('question_id') and str(q.get('question_type','')) not in {'TITLE','SECTION','DESCRIPTION'}: expected.append(str(q['question_id']))
        for c in expected:
            if c not in headers: errors.append(f'Missing configured data column: {c}')
    elif pt=='QUALITATIVE_INTERVIEW':
        for c in ['Sequence','Kind','Parent_Question_ID','Question_ID','Question','Answer']:
            if c not in headers: errors.append(f'Missing required column: {c}')
        valid={str(q.get('question_id')) for q in cfg.get('initial_questions',[])}
        for i,r in enumerate(rows,2):
            k=_s(r.get('Kind')).upper()
            if k not in {'INITIAL','FOLLOW_UP'}: errors.append(f'Row {i}: Kind must be INITIAL or FOLLOW_UP.')
            if k=='INITIAL' and _s(r.get('Question_ID')) not in valid: errors.append(f'Row {i}: unknown initial Question_ID.')
    else:
        z=bool(cfg.get('use_z_number'))
        for c in ['Row_Code','Column_Code','A_Code']+(['B_Code'] if z else []):
            if c not in headers: errors.append(f'Missing required column: {c}')
    for i,r in enumerate(rows,2):
        if not _s(r.get('Respondent_ID')): errors.append(f'Row {i}: Respondent_ID is required.')
    preview=rows[:50]
    return {'valid':not errors,'errors':errors,'warnings':warnings,'headers':headers,'rows':rows,'preview':preview,'project_type':pt,'mapping':mapping,'autodetect':bool(autodetect)}

def _ensure_respondent(db,r):
    rid=_s(r.get('Respondent_ID')); existing=db.get_respondent(rid)
    if existing:return rid
    username=_s(r.get('Username')) or ('import_'+rid)
    # keep import accounts non-login-oriented unless user later resets credentials
    try: return db.add_respondent(username=username,password=uuid.uuid4().hex,full_name=_s(r.get('Full_Name')) or rid,organization=_s(r.get('Organization')),respondent_id=rid)
    except Exception as e: raise ValueError(f"Cannot create respondent {rid}: {e}")

def import_file(db,path,autodetect=False):
    check=validate_import_file(db,path,autodetect=autodetect)
    if not check['valid']: raise ValueError('\n'.join(check['errors'][:20]))
    rows=check['rows']; pt=check['project_type']; cfg=db.get_config_dict(); imported=0
    grouped={}
    for r in rows: grouped.setdefault(_s(r.get('Respondent_ID')),[]).append(r)
    for rid,items in grouped.items():
        _ensure_respondent(db,items[0])
        if pt=='QUANTITATIVE_SURVEY':
            r=items[0]; data={}
            for q in cfg.get('questions',[]):
                qt=str(q.get('question_type','')); qid=str(q.get('question_id',''))
                if qt.startswith('MATRIX_'):
                    data[qid]={str(x.get('item_id')):r.get(str(x.get('item_id'))) for x in q.get('matrix_items',[]) if x.get('item_id')}
                elif qid and qt not in {'TITLE','SECTION','DESCRIPTION'}: data[qid]=r.get(qid)
        elif pt=='QUALITATIVE_INTERVIEW':
            data={'interview_items':[]}
            for r in items:
                data['interview_items'].append({'sequence':_s(r.get('Sequence')),'kind':_s(r.get('Kind')).upper(),'parent_question_id':_s(r.get('Parent_Question_ID')) or None,'question_id':_s(r.get('Question_ID')) or None,'question':_s(r.get('Question')),'answer':_s(r.get('Answer'))})
        else:
            # Long matrix import. Codes are preserved; numerical lookup is derived from current project scale when possible.
            factors=[x for x in cfg.get('factors',[]) if str(x.get('type','Factor')).lower()!='alternative']; alts=cfg.get('alternatives',[])
            model=str(cfg.get('model_type') or cfg.get('model') or '').upper(); row_codes=[_s(x.get('code')) for x in (alts if model=='TOPSIS' else factors)]; col_codes=['Rating'] if model=='DELPHI' else [_s(x.get('code')) for x in factors]
            if not row_codes: row_codes=sorted({_s(x.get('Row_Code')) for x in items})
            if not col_codes: col_codes=sorted({_s(x.get('Column_Code')) for x in items})
            idxr={x:i for i,x in enumerate(row_codes)}; idxc={x:i for i,x in enumerate(col_codes)}
            code=[['' for _ in col_codes] for _ in row_codes]; val=[['' for _ in col_codes] for _ in row_codes]
            z=bool(cfg.get('use_z_number')); za=[['' for _ in col_codes] for _ in row_codes]; zb=[['' for _ in col_codes] for _ in row_codes]
            amap={_s(x.get('code')):x.get('value','') for x in (cfg.get('z_scales_a',[]) if z else cfg.get('standard_scales',[]))}
            bmap={_s(x.get('code')):x.get('value','') for x in cfg.get('z_scales_b',[])}
            for r in items:
                rr=_s(r.get('Row_Code')); cc=_s(r.get('Column_Code'))
                if rr not in idxr or cc not in idxc: raise ValueError(f'{rid}: unknown matrix coordinate {rr}/{cc}.')
                a=_s(r.get('A_Code')); b=_s(r.get('B_Code'))
                if a and a not in amap: raise ValueError(f'{rid}: unknown A_Code {a}.')
                if z and b and b not in bmap: raise ValueError(f'{rid}: unknown B_Code {b}.')
                av=amap.get(a,''); bv=bmap.get(b,'')
                code[idxr[rr]][idxc[cc]]=f'{a}|{b}' if z else a
                val[idxr[rr]][idxc[cc]]=f'{av}|{bv}' if z else av
                za[idxr[rr]][idxc[cc]]=a; zb[idxr[rr]][idxc[cc]]=b
            data={'code_matrix':code,'value_matrix':val}
            if z:
                zav=[[amap.get(x,'') for x in row] for row in za]; zbv=[[bmap.get(x,'') for x in row] for row in zb]
                data.update({'z_a_code_matrix':za,'z_b_code_matrix':zb,'z_a_value_matrix':zav,'z_b_value_matrix':zbv})
        db.save_response(rid,data,source_file=f'LEGACY_IMPORT:{Path(path).name}'); imported+=1
    return {'responses':imported,'rows':len(rows),'file':str(path)}
