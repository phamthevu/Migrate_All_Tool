import json, tempfile
from pathlib import Path
import openpyxl
from core.project_db import ProjectDB
from core.models import ProjectConfig,Factor,ScaleItem,Expert,config_to_dict
from integrations.snapshot_staging import build_snapshot,split_z
from integrations.snapshot_package import create_snapshot_package
from integrations.s3_repository import S3Repository

def make_db(n=20,factors=6):
    td=tempfile.TemporaryDirectory(); path=Path(td.name)/'ztest.mcdm'
    cfg=ProjectConfig(model_type='DEMATEL',use_z_number=True,scale_type='Category',expert_auth_mode='Account',factors=[Factor('Factor',f'F{i}',f'Factor {i}') for i in range(factors)],z_scales_a=[ScaleItem('High','H','0.5, 0.75, 1.0'),ScaleItem('Very High','VH','0.75, 1, 1')],z_scales_b=[ScaleItem('Likely','Lk','0.5, 0.7, 0.9'),ScaleItem('Very Likely','VLk','0.8, 1, 1')],experts=[Expert(f'E{i}',f'u{i}','123',f'R{i}') for i in range(n)])
    accounts=[{'Respondent_ID':f'R{i}','Username':f'u{i}','Password':'123','Full_Name':f'E{i}'} for i in range(n)]
    ProjectDB.create_project(str(path),'Z Stress','MCDM',config_to_dict(cfg),accounts); db=ProjectDB(str(path))
    for i in range(n):
        code=[[('(H; Lk)' if (r+c+i)%2==0 else '(VH; VLk)') for c in range(factors)] for r in range(factors)]
        val=[[('(0.5, 0.75, 1.0; 0.5, 0.7, 0.9)' if (r+c+i)%2==0 else '(0.75, 1, 1; 0.8, 1, 1)') for c in range(factors)] for r in range(factors)]
        # explicit matrices are the V5.7.1 format
        za=[[split_z(x)[0] for x in row] for row in code]; zb=[[split_z(x)[1] for x in row] for row in code]
        # vector values contain commas; explicit A/B avoids ambiguous parsing
        zav=[['0.5, 0.75, 1.0' if x=='H' else '0.75, 1, 1' for x in row] for row in za]
        zbv=[['0.5, 0.7, 0.9' if x=='Lk' else '0.8, 1, 1' for x in row] for row in zb]
        db.save_response(f'R{i}',{'code_matrix':code,'value_matrix':val,'z_a_code_matrix':za,'z_b_code_matrix':zb,'z_a_value_matrix':zav,'z_b_value_matrix':zbv})
    return td,db

def test_z_persistence_flatten_and_package():
    td,db=make_db(20,6)
    try:
        row=db.get_response_by_respondent('R0'); p=json.loads(row['response_json']); assert p['z_b_code_matrix'][0][0]=='Lk'
        db.close(); db=ProjectDB(str(Path(td.name)/'ztest.mcdm')); p=json.loads(db.get_response_by_respondent('R0')['response_json']); assert p['z_b_code_matrix'][0][0]=='Lk'
        data,ident,sid,at,res,vals=build_snapshot(db); assert len(res)==20; assert len(vals)==20*36
        first=vals[0]; assert first['z_a_code']=='H' and first['z_b_code']=='Lk'; assert first['z_b_value']=='0.5, 0.7, 0.9'
        pkg=create_snapshot_package(db,td.name,stamp='20260926T000000Z'); assert len(pkg['files'])==3
        for f in pkg['files']: assert Path(f).is_file() and Path(f).stat().st_size>0
        wb=openpyxl.load_workbook(pkg['raw_excel'],read_only=True); assert 'Metadata' in wb.sheetnames and 'Ans_01' in wb.sheetnames; wb.close()
        wb=openpyxl.load_workbook(pkg['numerical_excel'],read_only=True); assert 'Ans_01' in wb.sheetnames; wb.close()
    finally:
        db.close(); td.cleanup()

def test_s3_uploads_exactly_three_files():
    td,db=make_db(2,3)
    class Fake:
        def __init__(self): self.calls=[]
        def put_object(self,**kw): self.calls.append(kw); return {}
    fake=Fake(); repo=S3Repository({'bucket':'b','prefix':'p'}); repo._client=lambda:fake
    try:
        out=repo.sync(db); assert out['count']==3; assert len(fake.calls)==3
        names=[x['Key'] for x in fake.calls]; assert any(x.endswith('.json') for x in names); assert any(x.endswith('_RAW_CODES.xlsx') for x in names); assert any(x.endswith('_NUMERICAL_VALUES.xlsx') for x in names)
    finally: db.close(); td.cleanup()
