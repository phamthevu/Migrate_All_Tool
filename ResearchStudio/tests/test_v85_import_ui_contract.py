from pathlib import Path
from openpyxl import load_workbook
from core.import_builder import validate_definition,generate_data_template

def test_quant_more_than_ten_questions(tmp_path):
    qs=[{'question_id':f'Q{i:03}','question_text':f'Question {i}','question_type':'SHORT_TEXT','required':False,'options':''} for i in range(1,26)]
    d={'project_type':'QUANTITATIVE_SURVEY','source_mode':'UI','access_mode':'Free','questions':qs,'matrix_items':[]}
    assert validate_definition(d)==[]
    p=tmp_path/'q.xlsx';generate_data_template(p,d)
    wb=load_workbook(p,read_only=True);assert wb['Questions'].max_row==26;assert wb['Data'].max_column==7+25;wb.close()

def test_matrix_requires_options_and_component():
    d={'project_type':'QUANTITATIVE_SURVEY','questions':[{'question_id':'Q001','question_text':'Matrix','question_type':'MATRIX_LIKERT','required':False,'options':''}],'matrix_items':[]}
    e=validate_definition(d)
    assert any('requires Options' in x for x in e)
    assert any('requires at least one matrix component' in x for x in e)

def test_matrix_valid_with_choice_scale_and_components():
    d={'project_type':'QUANTITATIVE_SURVEY','questions':[{'question_id':'Q001','question_text':'Matrix','question_type':'MATRIX_LIKERT','required':True,'options':'1|2|3|4|5'}],
       'matrix_items':[{'parent_id':'Q001','item_id':'Q001_01','item_text':'Quality'},{'parent_id':'Q001','item_id':'Q001_02','item_text':'Value'}]}
    assert validate_definition(d)==[]

def test_qual_more_than_ten_questions(tmp_path):
    qs=[{'question_id':f'Q{i:03}','question_text':f'Interview {i}','required':False} for i in range(1,31)]
    d={'project_type':'QUALITATIVE_INTERVIEW','source_mode':'UI','access_mode':'Free','questions':qs}
    assert validate_definition(d)==[]
    p=tmp_path/'qual.xlsx';generate_data_template(p,d)
    wb=load_workbook(p,read_only=True);assert wb['Questions'].max_row==31;wb.close()

def test_mcdm_more_than_ten_factors(tmp_path):
    factors=[{'code':f'C{i:03}','name':f'Factor {i}'} for i in range(1,31)]
    d={'project_type':'MCDM','source_mode':'UI','access_mode':'Free','model':'DEMATEL','use_z':False,
       'factors':factors,'alternatives':[],'scale':[{'term':'Low','code':'S01','value':'0,0,0.5'},{'term':'High','code':'S02','value':'0.5,1,1'}],'z_scale':[]}
    assert validate_definition(d)==[]
    p=tmp_path/'m.xlsx';generate_data_template(p,d)
    wb=load_workbook(p,read_only=True);assert wb['Factors_Alternatives'].max_row==31;assert wb['Ans_01'].max_row==31;wb.close()
