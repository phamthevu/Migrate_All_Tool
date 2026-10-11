from openpyxl import load_workbook, Workbook
from core.standalone_import import generate_standalone_template, inspect_workbook, detect_workbook, create_project_from_workbook
from core.project_db import ProjectDB

def test_quant_template_is_friendly_and_accounts_optional(tmp_path):
    x=tmp_path/'q.xlsx'; out=tmp_path/'q.mcdm'; generate_standalone_template(x,'QUANTITATIVE_SURVEY')
    wb=load_workbook(x)
    assert wb.sheetnames[0]=='START_HERE'
    assert wb['_System'].sheet_state=='hidden'
    assert wb['Questions'].sheet_properties.tabColor is not None
    assert wb['Data'].sheet_properties.tabColor is not None
    assert wb['Accounts'].sheet_properties.tabColor is not None
    # Free mode does not require account rows.
    ws=wb['Accounts']; ws.delete_rows(2, ws.max_row); wb.save(x)
    assert inspect_workbook(x)['valid']
    r=create_project_from_workbook(x,out,'No Accounts')
    db=ProjectDB(out); assert db.get_meta()['project_id'].startswith('PRJ_'); db.close()

def test_ordinary_excel_is_scannable_but_invalid(tmp_path):
    x=tmp_path/'ordinary.xlsx'; wb=Workbook(); ws=wb.active; ws.title='Sheet1'; ws.append(['name','score']); ws.append(['A',5]); wb.save(x)
    scan=detect_workbook(x); assert scan['project_type']=='UNKNOWN' and 'Sheet1' in scan['sheets']
    c=inspect_workbook(x); assert not c['valid']; assert any('not a Research Studio import template' in e for e in c['errors'])

def test_quant_template_has_no_user_project_identity_fields(tmp_path):
    x=tmp_path/'q.xlsx'; generate_standalone_template(x,'QUANTITATIVE_SURVEY'); wb=load_workbook(x,data_only=True)
    setup=[r[0] for r in wb['Project_Setup'].iter_rows(min_row=2,values_only=True)]
    assert 'Project ID' not in setup and 'Storage Key' not in setup
    headers=[c.value for c in wb['Data'][1]]
    assert 'Source_Project_ID' not in headers and 'Source_Project' not in headers and 'Source_File' not in headers


def test_account_mode_requires_account_row(tmp_path):
    x=tmp_path/'q_account.xlsx'; generate_standalone_template(x,'QUANTITATIVE_SURVEY'); wb=load_workbook(x)
    wb['Project_Setup']['B3']='Account'; wb.save(x)
    c=inspect_workbook(x); assert not c['valid']; assert any('requires at least one account' in e for e in c['errors'])
