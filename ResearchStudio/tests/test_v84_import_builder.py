from pathlib import Path
from openpyxl import Workbook, load_workbook
from core.import_builder import validate_definition, infer_from_sample, generate_data_template
from core.standalone_import import inspect_workbook, create_project_from_workbook
from core.project_db import ProjectDB

def test_v84_quant_ui_template_to_project(tmp_path):
    d={'project_type':'QUANTITATIVE_SURVEY','source_mode':'UI','access_mode':'Free','questions':[{'question_id':'AGE','question_text':'Age','question_type':'INTEGER','required':True,'options':''},{'question_id':'SAT','question_text':'Satisfaction','question_type':'MATRIX_LIKERT','required':True,'options':'1|2|3|4|5'}],'matrix_items':[{'parent_id':'SAT','item_id':'SAT1','item_text':'Quality'},{'parent_id':'SAT','item_id':'SAT2','item_text':'Value'}]}
    assert validate_definition(d)==[]
    x=tmp_path/'q.xlsx';generate_data_template(x,d)
    wb=load_workbook(x);ws=wb['Data'];ws.delete_rows(2);ws.append(['Ans_001','R001','A','','','','',30,4,5]);wb.save(x)
    c=inspect_workbook(x);assert c['valid'];assert c['project_type']=='QUANTITATIVE_SURVEY';assert len(c['preview'])==1
    p=tmp_path/'q.mcdm';res=create_project_from_workbook(x,p);assert res['responses']==1
    db=ProjectDB(p);assert db.conn.execute('SELECT COUNT(*) FROM responses').fetchone()[0]==1;db.close()

def test_v84_sample_quant_requires_data_row(tmp_path):
    x=tmp_path/'raw.xlsx';wb=Workbook();ws=wb.active;ws.title='Raw';ws.append(['AGE','SCORE']);wb.save(x)
    try:infer_from_sample(x,'QUANTITATIVE_SURVEY')
    except ValueError as e:assert 'at least one data row' in str(e)
    else:assert False

def test_v84_sample_quant_infers_columns(tmp_path):
    x=tmp_path/'raw.xlsx';wb=Workbook();ws=wb.active;ws.title='Raw';ws.append(['Respondent_ID','AGE','SCORE','NOTE']);ws.append(['R1',28,4,'ok']);wb.save(x)
    d=infer_from_sample(x,'QUANTITATIVE_SURVEY');assert [q['question_id'] for q in d['questions']]==['AGE','SCORE','NOTE']

def test_v84_mcdm_sample_requires_ans(tmp_path):
    x=tmp_path/'raw.xlsx';wb=Workbook();wb.active.title='Whatever';wb.save(x)
    try:infer_from_sample(x,'MCDM')
    except ValueError as e:assert 'ANS' in str(e)
    else:assert False

def test_v84_mcdm_detect_and_generate(tmp_path):
    x=tmp_path/'raw.xlsx';wb=Workbook();ws=wb.active;ws.title='ANS_OLD';ws.append(['','C1','C2']);ws.append(['C1','L','H']);ws.append(['C2','M','L']);wb.save(x)
    d=infer_from_sample(x,'MCDM','DEMATEL',False);assert len(d['factors'])==2
    out=tmp_path/'m.xlsx';generate_data_template(out,d);w=load_workbook(out);assert 'Ans_01' in w.sheetnames;assert w['_System'].sheet_state=='hidden'

def test_v84_qual_sample_requires_answer(tmp_path):
    x=tmp_path/'raw.xlsx';wb=Workbook();ws=wb.active;ws.append(['Question_ID','Question']);ws.append(['Q1','Hello']);wb.save(x)
    try:infer_from_sample(x,'QUALITATIVE_INTERVIEW')
    except ValueError as e:assert 'Answer' in str(e)
    else:assert False
