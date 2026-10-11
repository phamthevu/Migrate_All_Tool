from pathlib import Path
from openpyxl import load_workbook
from core.standalone_import import generate_standalone_template, inspect_workbook, create_project_from_workbook
from core.project_db import ProjectDB
from core.survey_models import survey_config_to_dict, SurveyQuestion


def test_quant_template_preview_creates_new_project(tmp_path):
    x=tmp_path/'q.xlsx'; out=tmp_path/'q.mcdm'
    generate_standalone_template(x,'QUANTITATIVE_SURVEY')
    c=inspect_workbook(x); assert c['valid'] and c['sheet']=='Data' and c['preview']
    r=create_project_from_workbook(x,out,'Imported Quant')
    db=ProjectDB(out); m=db.get_meta(); rows=db.get_all_responses()
    assert r['responses']==1 and len(rows)==1
    assert m['project_id'].startswith('PRJ_') and m['storage_key']
    assert rows[0]['source_file']=='q.xlsx'
    db.close()

def test_qual_template_creates_new_project(tmp_path):
    x=tmp_path/'qual.xlsx'; out=tmp_path/'qual.mcdm'
    generate_standalone_template(x,'QUALITATIVE_INTERVIEW')
    assert inspect_workbook(x)['valid']
    r=create_project_from_workbook(x,out,'Imported Qual')
    assert r['responses']==1

def test_mcdm_template_is_export_like_multisheet_and_imports(tmp_path):
    x=tmp_path/'m.xlsx'; out=tmp_path/'m.mcdm'
    generate_standalone_template(x,'MCDM','DEMATEL',False)
    wb=load_workbook(x,read_only=True); assert {'START_HERE','Project_Setup','_System','Lookup','Ans_01','Factors_Alternatives','Scale_Standard','Schema'} <= set(wb.sheetnames)
    assert inspect_workbook(x)['valid']
    r=create_project_from_workbook(x,out,'Imported MCDM'); assert r['responses']==1

def test_invalid_workbook_blocked(tmp_path):
    x=tmp_path/'bad.xlsx'; generate_standalone_template(x,'QUANTITATIVE_SURVEY')
    wb=load_workbook(x); del wb['Data']; wb.save(x)
    c=inspect_workbook(x); assert not c['valid'] and any('Data' in e for e in c['errors'])

def test_import_twice_produces_distinct_project_identity(tmp_path):
    x=tmp_path/'q.xlsx'; generate_standalone_template(x,'QUANTITATIVE_SURVEY')
    a=tmp_path/'a.mcdm'; b=tmp_path/'b.mcdm'
    create_project_from_workbook(x,a,'Same Name'); create_project_from_workbook(x,b,'Same Name')
    da,db=ProjectDB(a),ProjectDB(b)
    assert da.get_meta()['project_id'] != db.get_meta()['project_id']
    assert da.get_meta()['storage_key'] != db.get_meta()['storage_key']
    da.close();db.close()

def test_merge_free_and_legacy_anonymous_compatible_without_password_prompt_semantics(tmp_path):
    q=[SurveyQuestion('Q1','Question','SHORT_TEXT',False)]
    a,b=tmp_path/'a.mcdm',tmp_path/'b.mcdm'
    ProjectDB.create_project(a,'Free','QUANTITATIVE_SURVEY',survey_config_to_dict(q,'Free'),[])
    # legacy naming: merge compatibility must ignore access/auth mode entirely
    cfg=survey_config_to_dict(q,'Anonymous')
    ProjectDB.create_project(b,'Anonymous','QUANTITATIVE_SURVEY',cfg,[])
    s=ProjectDB(b); rid=s.create_free_respondent('Legacy Anonymous'); s.save_response(rid,{'Q1':'old'}); s.close()
    t=ProjectDB(a); assert t.compare_merge_config(b)['compatible']; t.merge_other_projects([b])
    rows=t.get_all_responses(); assert len(rows)==1 and rows[0]['source_project_id']
    # Data Manager editing uses respondent identity directly; merge does not transform a free response into a login challenge.
    assert t.get_response_by_respondent(rows[0]['respondent_id']) is not None
    t.close()
