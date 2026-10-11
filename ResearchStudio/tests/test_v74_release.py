import json,tempfile
from pathlib import Path
import openpyxl
from core.project_db import ProjectDB
from core.project_excel import export_quantitative,export_qualitative
from integrations.snapshot_package import create_snapshot_package

def test_quant_ui_has_no_followup_and_no_implicit_comment():
 s=Path('ui/tab_quant_survey.py').read_text(encoding='utf-8')
 assert 'Add follow-up' not in s and 'def add_followup' not in s
 assert "QGroupBox('Comment')" not in s and "data[\"_comment\"]" not in s

def test_quant_export_is_one_flat_excel_without_implicit_comment():
 with tempfile.TemporaryDirectory() as td:
  p=Path(td)/'q.mcdm';cfg={'auth_mode':'Free','questions':[{'question_id':'Q1','question_text':'Age','question_type':'INTEGER','required':False,'display_order':1},{'question_id':'Q2','question_text':'Tags','question_type':'MULTIPLE_CHOICE','required':False,'options':['A','B'],'display_order':2}]};ProjectDB.create_project(str(p),'Survey','QUANTITATIVE_SURVEY',cfg,[]);db=ProjectDB(str(p));rid=db.create_free_respondent('P1');db.save_response(rid,{'Q1':25,'Q2':['A','B'],'_comment':'legacy ignored'});out=Path(td)/'q.xlsx';export_quantitative(db,out);wb=openpyxl.load_workbook(out);ws=wb['Data'];hdr=[c.value for c in ws[1]];assert 'Comment' not in hdr;assert ws.cell(2,hdr.index('Q1')+1).value==25;pkg=create_snapshot_package(db,Path(td)/'snap',stamp='T');assert len(pkg['excel_files'])==1 and len(pkg['files'])==2 and pkg['excel_files'][0].endswith('_QUANTITATIVE.xlsx');db.close()

def test_qual_snapshot_is_json_plus_one_excel():
 with tempfile.TemporaryDirectory() as td:
  p=Path(td)/'i.mcdm';cfg={'auth_mode':'Free','initial_questions':[{'question_id':'I1','question_text':'Why?','required':False,'display_order':1}]};ProjectDB.create_project(str(p),'Interview','QUALITATIVE_INTERVIEW',cfg,[]);db=ProjectDB(str(p));rid=db.create_free_respondent('P1');db.save_response(rid,{'interview_items':[{'sequence':'1','kind':'INITIAL','question_id':'I1','question':'Why?','answer':'Because'}]});pkg=create_snapshot_package(db,Path(td)/'snap',stamp='T');assert len(pkg['excel_files'])==1 and len(pkg['files'])==2 and pkg['excel_files'][0].endswith('_QUALITATIVE.xlsx');wb=openpyxl.load_workbook(pkg['excel_files'][0]);assert 'Data' in wb.sheetnames;db.close()

def test_success_and_error_share_friendly_dialog_structure():
 s=Path('ui/friendly_message.py').read_text(encoding='utf-8');assert 'What happened' in s and 'Recommended action / What you can do:' in s and 'Technical details' in s and 'Copy Details' in s

def test_email_text_is_project_type_aware():
 s=Path('ui/dialog_file_transfer.py').read_text(encoding='utf-8');assert 'MCDM = JSON + 2 Excel files' in s and 'Quantitative = JSON + 1 Excel file' in s and 'Qualitative = JSON + 1 Excel file' in s
