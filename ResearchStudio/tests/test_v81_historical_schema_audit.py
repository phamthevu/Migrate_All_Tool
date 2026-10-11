import os,tempfile,time
from openpyxl import load_workbook
from core.project_db import ProjectDB
from core.legacy_data_import import generate_import_template,validate_import_file,import_file
from core.project_excel import export_quantitative

def cfg():
 return {'survey_schema_version':2,'auth_mode':'Free','questions':[{'question_id':'AGE','question_text':'Age','question_type':'INTEGER','required':False,'matrix_items':[]},{'question_id':'SAT','question_text':'Sat','question_type':'MATRIX_LIKERT','required':False,'matrix_items':[{'item_id':'SAT1','item_text':'One','display_order':1}]}]}

def test_v81_template_has_schema_and_preview():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm');x=os.path.join(d,'i.xlsx');ProjectDB.create_project(p,'Q','QUANTITATIVE_SURVEY',cfg(),[]);db=ProjectDB(p)
  generate_import_template(db,x); wb=load_workbook(x); assert {'Data','Schema','README'} <= set(wb.sheetnames)
  c=validate_import_file(db,x); assert c['valid'] and c['preview']

def test_v81_autodetect_schema_source_column_mapping():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm');x=os.path.join(d,'i.xlsx');ProjectDB.create_project(p,'Q','QUANTITATIVE_SURVEY',cfg(),[]);db=ProjectDB(p)
  generate_import_template(db,x);wb=load_workbook(x);ws=wb['Data']; headers=[c.value for c in ws[1]]; age=headers.index('AGE')+1; ws.cell(1,age).value='legacy_age'
  sc=wb['Schema']
  for row in sc.iter_rows(min_row=2):
   if row[0].value=='AGE': row[4].value='legacy_age'
  wb.save(x)
  strict=validate_import_file(db,x,False); assert not strict['valid']
  auto=validate_import_file(db,x,True); assert auto['valid'],auto['errors']; assert auto['mapping']['AGE']=='legacy_age'

def test_v81_created_preserved_modified_changes_and_source_audit():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm');x=os.path.join(d,'old_file.xlsx');ProjectDB.create_project(p,'Q','QUANTITATIVE_SURVEY',cfg(),[]);db=ProjectDB(p)
  generate_import_template(db,x);wb=load_workbook(x);ws=wb['Data'];ws.delete_rows(2);ws.append(['R1','u1','A','Org',20,4]);wb.save(x);import_file(db,x)
  r1=db.get_all_responses()[0]; created=r1['submitted_at']; modified=r1['modified_at']; assert r1['source_project_id']==db.get_meta()['project_id']; assert 'old_file.xlsx' in r1['source_file']; assert created==modified
  time.sleep(1.05); db.save_response('R1',{'AGE':21,'SAT':{'SAT1':5}},source_file=r1['source_file']);r2=db.get_all_responses()[0]
  assert r2['submitted_at']==created; assert r2['modified_at']>modified

def test_v81_quant_export_has_schema_and_audit_columns():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm');x=os.path.join(d,'out.xlsx');ProjectDB.create_project(p,'Q','QUANTITATIVE_SURVEY',cfg(),[]);db=ProjectDB(p);db.add_respondent('u','pw','A','',respondent_id='R1');db.save_response('R1',{'AGE':22,'SAT':{'SAT1':5}});export_quantitative(db,x)
  wb=load_workbook(x,read_only=True); assert 'Schema' in wb.sheetnames; h=[c.value for c in next(wb['Data'].iter_rows())]; assert 'Source_Project_ID' in h and 'Source_File' in h and 'Modified_At' in h
