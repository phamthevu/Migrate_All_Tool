import os,tempfile
from openpyxl import load_workbook
from core.project_db import ProjectDB
from core.legacy_data_import import generate_import_template,validate_import_file,import_file

def _quant_cfg():
 return {'survey_schema_version':2,'auth_mode':'Free','questions':[{'question_id':'AGE','question_text':'Age','question_type':'INTEGER','required':False,'matrix_items':[]},{'question_id':'SAT','question_text':'Sat','question_type':'MATRIX_LIKERT','required':False,'matrix_items':[{'item_id':'SAT1','item_text':'One','display_order':1}]}]}

def test_quant_template_validate_import():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm'); x=os.path.join(d,'import.xlsx'); ProjectDB.create_project(p,'Legacy Quant','QUANTITATIVE_SURVEY',_quant_cfg(),[]); db=ProjectDB(p)
  generate_import_template(db,x); wb=load_workbook(x); ws=wb['Data']; ws.delete_rows(2); ws.append(['R1','u1','Legacy A','Org',30,5]); wb.save(x)
  c=validate_import_file(db,x); assert c['valid'],c['errors']; r=import_file(db,x); assert r['responses']==1; rows=db.get_all_responses(); assert len(rows)==1 and rows[0]['respondent_id']=='R1'

def test_wrong_template_reports_missing_columns():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm'); x=os.path.join(d,'bad.xlsx'); ProjectDB.create_project(p,'Legacy Quant','QUANTITATIVE_SURVEY',_quant_cfg(),[]); db=ProjectDB(p)
  generate_import_template(db,x); wb=load_workbook(x); ws=wb['Data']; ws.delete_cols(1); wb.save(x)
  c=validate_import_file(db,x); assert not c['valid']; assert any('Respondent_ID' in e for e in c['errors'])

def test_config_hash_unchanged_after_import():
 with tempfile.TemporaryDirectory() as d:
  p=os.path.join(d,'p.mcdm'); x=os.path.join(d,'import.xlsx'); ProjectDB.create_project(p,'Legacy Quant','QUANTITATIVE_SURVEY',_quant_cfg(),[]); db=ProjectDB(p); before=db.get_meta()['config_hash']
  generate_import_template(db,x); wb=load_workbook(x); ws=wb['Data']; ws.delete_rows(2); ws.append(['R1','u1','A','',22,4]); wb.save(x); import_file(db,x)
  assert db.get_meta()['config_hash']==before
