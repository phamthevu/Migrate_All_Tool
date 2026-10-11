
import json
from openpyxl import load_workbook
from core.import_builder import generate_data_template
from core.standalone_import import inspect_workbook
from core.project_db import ProjectDB

def _quant_def():
    return {
        'project_type':'QUANTITATIVE_SURVEY','source_mode':'UI','access_mode':'Free',
        'questions':[
            {'question_id':'Q001','question_text':'Age','question_type':'INTEGER','required':False,'options':''},
            {'question_id':'Q002','question_text':'Satisfaction','question_type':'MATRIX_LIKERT','required':True,'options':'1|2|3|4|5'},
        ],
        'matrix_items':[
            {'parent_id':'Q002','item_id':'Q002_01','item_text':'Quality'},
            {'parent_id':'Q002','item_id':'Q002_02','item_text':'Value'},
        ],
    }

def test_generated_quant_schema_passes(tmp_path):
    p=tmp_path/'good.xlsx'; generate_data_template(p,_quant_def())
    c=inspect_workbook(p)
    assert c['valid'], c['errors']

def test_wrong_quant_data_schema_is_blocked(tmp_path):
    p=tmp_path/'bad.xlsx'; generate_data_template(p,_quant_def())
    wb=load_workbook(p); ws=wb['Data']; ws.cell(1,8).value='WRONG_COLUMN'; wb.save(p); wb.close()
    c=inspect_workbook(p)
    assert not c['valid']
    assert any('missing configured' in e.lower() or 'not defined' in e.lower() or 'headers/order' in e.lower() for e in c['errors'])

def test_choice_without_options_is_blocked_at_preview(tmp_path):
    d=_quant_def(); d['questions'][1]['options']=''
    # bypass builder validation by first making valid file then corrupt Questions
    p=tmp_path/'bad_options.xlsx'; generate_data_template(p,_quant_def())
    wb=load_workbook(p); ws=wb['Questions']
    for row in range(2,ws.max_row+1):
        if ws.cell(row,1).value=='Q002': ws.cell(row,5).value=''
    wb.save(p); wb.close()
    c=inspect_workbook(p)
    assert not c['valid']
    assert any('requires Options' in e for e in c['errors'])

def test_get_response_by_id_returns_exact_row(tmp_path):
    p=tmp_path/'x.mcdm'
    cfg={'questions':[]}
    ProjectDB.create_project(str(p),'x','QUANTITATIVE_SURVEY',cfg,[])
    db=ProjectDB(str(p))
    try:
        rid1=db.create_free_respondent('A'); rid2=db.create_free_respondent('B')
        db.save_response(rid1,{'v':'first'}); db.save_response(rid2,{'v':'second'})
        rows=db.get_all_responses()
        r2=db.get_response_by_id(rows[1]['id'])
        assert r2['respondent_id']==rid2
        assert json.loads(r2['response_json'])['v']=='second'
    finally:
        db.close()
